"""
Milestone 4 tests: PDF/Excel report generation, empty datasets,
permissions and data accuracy. Uses a throwaway SQLite DB (conftest.py).
Run from the backend folder:  python -m pytest tests -v
"""
import io
import uuid
from datetime import datetime, timedelta, timezone

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.database import Base, SessionLocal, engine
from app.main import app
from app.models.freshness_assessment import FreshnessAssessment, FreshnessCategory
from app.models.storage_condition import ComplianceStatus, StorageCondition

client = TestClient(app)
TODAY = datetime.now(timezone.utc).date()
INVENTORY_KINDS = ["freshness", "shelf-life", "inventory-quality", "waste-reduction", "storage-compliance"]


@pytest.fixture(autouse=True)
def fresh_db():
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)
    yield


def make_user(role, name="Test User"):
    email, pw = f"{uuid.uuid4().hex[:8]}@example.com", "Passw0rdTest1"
    r = client.post("/api/auth/register", json={
        "full_name": name, "email": email, "phone": "+911234567890",
        "password": pw, "confirm_password": pw, "role": role})
    assert r.status_code == 201, r.text
    t = client.post("/api/auth/login", json={"email": email, "password": pw}).json()
    return r.json()["id"], {"Authorization": f"Bearer {t['access_token']}"}, email


def add_item(h, name="Apple", category="Fruits", batch="apl-002", qty=5, days_to_expiry=20):
    r = client.post("/api/inventory", headers=h, json={
        "food_name": name, "category": category, "batch_number": batch, "quantity": qty,
        "purchase_date": str(TODAY - timedelta(days=10)),
        "expiry_date": str(TODAY + timedelta(days=days_to_expiry)), "storage_location": "fridge"})
    assert r.status_code == 201, r.text
    return r.json()["id"]


def add_assessment(user_id, item_id, score=88, cat=FreshnessCategory.GOOD):
    db = SessionLocal()
    db.add(FreshnessAssessment(
        inventory_item_id=uuid.UUID(item_id), user_id=uuid.UUID(user_id), image_path="x.jpg",
        freshness_score=score, freshness_category=cat, spoilage_probability=0.12,
        color_analysis={}, texture_analysis={}, mold_detected=False, bruising_detected=True,
        physical_damage_detected=False, quality_assessment="Item appears in good condition"))
    db.commit(); db.close()


def add_storage(user_id, item_id, temp=3.0, hum=85.0):
    db = SessionLocal()
    db.add(StorageCondition(
        inventory_item_id=uuid.UUID(item_id), user_id=uuid.UUID(user_id), temperature_celsius=temp,
        humidity_percent=hum, air_circulation="Good", light_exposure="Low", packaging_type="Box",
        compliance_status=ComplianceStatus.SUITABLE, compliance_notes="Within range"))
    db.commit(); db.close()


def get(kind, fmt, headers):
    return client.get(f"/api/reports/{kind}?format={fmt}", headers=headers)


def sheet_rows(resp, index):
    wb = load_workbook(io.BytesIO(resp.content))
    return [[c for c in row] for row in wb.worksheets[index].iter_rows(values_only=True)]


def all_text(resp):
    wb = load_workbook(io.BytesIO(resp.content))
    return " ".join(str(c) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for c in row if c is not None)


# ---------------- generation ----------------
@pytest.mark.parametrize("kind", INVENTORY_KINDS)
def test_empty_dataset_generates_valid_files(kind):
    _, h, _ = make_user("consumer")
    pdf = get(kind, "pdf", h)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    assert "attachment" in pdf.headers["content-disposition"]
    x = get(kind, "xlsx", h)
    assert x.status_code == 200 and x.content[:2] == b"PK"
    assert load_workbook(io.BytesIO(x.content)).sheetnames[0] == "Summary"
    assert "No " in all_text(x) or "Not " in all_text(x)


@pytest.mark.parametrize("kind", INVENTORY_KINDS)
def test_populated_dataset_generates_valid_files(kind):
    uid, h, _ = make_user("consumer")
    iid = add_item(h, days_to_expiry=3); add_assessment(uid, iid); add_storage(uid, iid)
    assert get(kind, "pdf", h).content.startswith(b"%PDF")
    assert "Apple" in all_text(get(kind, "xlsx", h))


# ---------------- data accuracy ----------------
def test_freshness_values_match_stored_records_and_unassessed_is_not_zero():
    uid, h, _ = make_user("consumer")
    a = add_item(h, "Apple", batch="A1"); add_item(h, "Milk", "Dairy Products", "M1")
    add_assessment(uid, a, score=88)
    rows = sheet_rows(get("freshness", "xlsx", h), 1)
    hdr = rows[0]
    by = {r[0]: dict(zip(hdr, r)) for r in rows[1:]}
    assert by["Apple"]["Freshness score (0-100)"] == 88
    assert by["Apple"]["Spoilage probability (%)"] == 12.0
    assert by["Apple"]["Bruising flag"] == "Yes" and by["Apple"]["Mold flag"] == "No"
    assert by["Milk"]["Freshness score (0-100)"] == "Not assessed"
    assert by["Milk"]["Spoilage probability (%)"] == "Not assessed"


def test_storage_report_real_readings_and_missing_readings():
    uid, h, _ = make_user("consumer")
    a = add_item(h, "Apple"); add_item(h, "Milk", "Dairy Products", "M1")
    add_storage(uid, a, temp=3.0, hum=85.0)
    rows = sheet_rows(get("storage-compliance", "xlsx", h), 1)
    by = {r[0]: dict(zip(rows[0], r)) for r in rows[1:]}
    assert by["Apple"]["Temperature (C)"] == 3.0 and by["Apple"]["Humidity (%)"] == 85.0
    assert by["Apple"]["Compliance"] == "Suitable"
    assert by["Milk"]["Temperature (C)"] == "No readings recorded"
    summary = dict((r[0], r[1]) for r in sheet_rows(get("storage-compliance", "xlsx", h), 0) if r[0])
    assert summary["Items with a storage reading"] == 1 and summary["Items with no reading"] == 1
    assert summary["Compliance rate (latest readings, Suitable / items with readings)"] == "100%"


def test_storage_compliance_rate_unavailable_without_readings():
    _, h, _ = make_user("consumer"); add_item(h)
    summary = dict((r[0], r[1]) for r in sheet_rows(get("storage-compliance", "xlsx", h), 0) if r[0])
    assert summary["Compliance rate (latest readings, Suitable / items with readings)"] == "Not available"


def test_shelf_life_matches_analysis_endpoint():
    _, h, _ = make_user("consumer")
    iid = add_item(h, days_to_expiry=10)
    api = client.get(f"/api/inventory/{iid}/analysis", headers=h).json()["shelf_life"]["current"]
    rows = sheet_rows(get("shelf-life", "xlsx", h), 1)
    row = dict(zip(rows[0], rows[1]))
    assert row["Remaining days (estimate)"] == api["remaining_shelf_life_days"]
    assert row["Risk"] == api["risk_level"]
    assert row["Forecast expiry"] == api["forecasted_expiry_date"]


def test_waste_report_lists_only_at_risk_items():
    _, h, _ = make_user("consumer")
    add_item(h, "Safe", "Packaged Foods", days_to_expiry=60, batch="S1"); add_item(h, "Soon", days_to_expiry=2, batch="N1")
    add_item(h, "Gone", days_to_expiry=-3, batch="G1", qty=7)
    rows = sheet_rows(get("waste-reduction", "xlsx", h), 1)
    names = {r[0] for r in rows[1:]}
    assert names == {"Soon", "Gone"}
    summary = dict((r[0], r[1]) for r in sheet_rows(get("waste-reduction", "xlsx", h), 0) if r[0])
    assert summary["Units already past expiry"] == 7


def test_waste_report_empty_message_when_nothing_at_risk():
    _, h, _ = make_user("consumer"); add_item(h, "Crackers", "Packaged Foods", days_to_expiry=60)
    assert "No items are currently at risk of waste." in all_text(get("waste-reduction", "xlsx", h))


def test_inventory_quality_discloses_neutral_defaults():
    _, h, _ = make_user("consumer"); add_item(h)
    rows = sheet_rows(get("inventory-quality", "xlsx", h), 1)
    row = dict(zip(rows[0], rows[1]))
    assert "Neutral default used for: visual, storage" in row["Health score data basis"]
    assert row["Freshness score"] == "Not assessed"
    summary = dict((r[0], r[1]) for r in sheet_rows(get("inventory-quality", "xlsx", h), 0) if r[0])
    assert summary["Average freshness score (assessed items)"] == "Not available"


def test_spreadsheet_formula_injection_is_neutralised():
    _, h, _ = make_user("consumer"); add_item(h, name="=HYPERLINK(\"http://x\")")
    rows = sheet_rows(get("freshness", "xlsx", h), 1)
    assert rows[1][0].startswith("'=")


# ---------------- permissions ----------------
@pytest.mark.parametrize("kind", INVENTORY_KINDS + ["platform-summary"])
def test_unauthenticated_rejected(kind):
    assert client.get(f"/api/reports/{kind}?format=pdf").status_code in (401, 403)


def test_invalid_token_rejected():
    assert get("freshness", "pdf", {"Authorization": "Bearer nope"}).status_code == 401


@pytest.mark.parametrize("kind", INVENTORY_KINDS)
def test_administrator_cannot_download_inventory_reports(kind):
    _, h, _ = make_user("administrator")
    assert get(kind, "pdf", h).status_code == 403


@pytest.mark.parametrize("role", ["consumer", "retail_manager", "warehouse_operator", "food_quality_inspector"])
def test_non_admin_cannot_download_platform_summary(role):
    _, h, _ = make_user(role)
    assert get("platform-summary", "pdf", h).status_code == 403


def test_unknown_kind_and_bad_format():
    _, h, _ = make_user("consumer")
    assert get("nope", "pdf", h).status_code == 404
    assert get("freshness", "docx", h).status_code == 422


# ---------------- shared inventory consistency ----------------
def test_shared_roles_see_identical_rows_and_consumer_is_isolated():
    ruid, ret, _ = make_user("retail_manager")
    _, wh, _ = make_user("warehouse_operator")
    _, ins, _ = make_user("food_quality_inspector")
    _, con, _ = make_user("consumer")
    iid = add_item(ret, "Apple", days_to_expiry=3); add_item(ret, "Milk", "Dairy Products", "M1", days_to_expiry=3)
    add_assessment(ruid, iid, score=91, cat=FreshnessCategory.FRESH)
    for kind in INVENTORY_KINDS:
        snaps = [sheet_rows(get(kind, "xlsx", h), 1) for h in (ret, wh, ins)]
        assert snaps[0] == snaps[1] == snaps[2], kind
        assert len(snaps[0]) >= 3, kind  # header + 2 items (waste may differ below)
    assert sheet_rows(get("freshness", "xlsx", con), 1)[1][0] == "No inventory items found."
    assert "Apple" not in all_text(get("freshness", "xlsx", con))


# ---------------- platform summary ----------------
def test_platform_summary_for_admin_is_aggregate_and_has_no_personal_data():
    _, adm, admin_email = make_user("administrator", name="Alice Admin")
    cuid, con, con_email = make_user("consumer", name="Zed Consumer")
    iid = add_item(con); add_assessment(cuid, iid); add_storage(cuid, iid)
    x = get("platform-summary", "xlsx", adm)
    assert x.status_code == 200
    text = all_text(x)
    for pii in (con_email, admin_email, "Zed Consumer", "+911234567890"):
        assert pii not in text
    summary = dict((r[0], r[1]) for r in sheet_rows(x, 0) if r[0])
    assert summary["Registered users"] == 2 and summary["Inventory items"] == 1
    assert summary["Freshness assessments"] == 1 and summary["Storage readings"] == 1
    assert get("platform-summary", "pdf", adm).content.startswith(b"%PDF")


def test_platform_summary_empty_platform():
    _, adm, _ = make_user("administrator")
    summary = dict((r[0], r[1]) for r in sheet_rows(get("platform-summary", "xlsx", adm), 0) if r[0])
    assert summary["Inventory items"] == 0 and summary["Freshness assessments"] == 0
