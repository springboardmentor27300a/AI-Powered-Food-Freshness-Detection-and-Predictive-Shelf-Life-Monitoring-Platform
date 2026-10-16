import io
import json
import os
from datetime import date
from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from database import SessionLocal
from models.batch import Batch
from models.food_item import FoodItem
from models.prediction_history import PredictionHistory
from models.storage_reading import StorageReading
from models.freshness_assessment import FreshnessAssessment
from services.storage_monitoring import analyze_storage_conditions
from services.inventory_rotation import generate_rotation_plan

router = APIRouter(prefix="/reports", tags=["Reports & Export"])


def get_food_map(db):
    return {food.id: food.name for food in db.query(FoodItem).all()}


def freshness_status(score):
    if score is None:
        return "Not Assessed"
    if score >= 80:
        return "Fresh"
    if score >= 65:
        return "Good"
    if score >= 50:
        return "Acceptable"
    if score >= 40:
        return "Near Spoilage"
    return "Spoiled"


def expiry_status(days):
    if days < 0:
        return "Expired"
    if days <= 3:
        return "Critical"
    if days <= 7:
        return "Near Expiry"
    return "Fresh"


def _fefo_entry_for_batch(db, batch_id):
    if not batch_id:
        return None
    batches = (
        db.query(Batch)
        .filter(Batch.quantity > 0)
        .all()
    )
    plan = generate_rotation_plan(batches)
    for item in plan:
        if item["batch_id"] == batch_id:
            return {
                "rotation_method": "FEFO - First Expired, First Out",
                "priority_rank": item.get("rotation_priority"),
                "effective_expiry_date": str(item.get("effective_expiry_date")) if item.get("effective_expiry_date") else None,
                "days_until_expiry": item.get("days_until_expiry"),
                "rotation_action": item.get("rotation_action"),
                "availability": item.get("availability"),
            }
    return {
        "rotation_method": "FEFO - First Expired, First Out",
        "priority_rank": None,
        "effective_expiry_date": None,
        "days_until_expiry": None,
        "rotation_action": "Batch is out of stock or unavailable for active FEFO rotation.",
        "availability": "Out of Stock",
    }


def _freshness_report(db):
    batches = db.query(Batch).all()
    food_map = get_food_map(db)
    records = []
    scores = []
    for batch in batches:
        if batch.freshness_score is not None:
            scores.append(batch.freshness_score)
        records.append({
            "batch_id": batch.id,
            "food_name": food_map.get(batch.food_item_id, "Unknown Food"),
            "quantity": batch.quantity,
            "freshness_score": batch.freshness_score,
            "freshness_status": freshness_status(batch.freshness_score),
            "spoilage_probability": batch.spoilage_probability,
            "spoilage_risk": batch.spoilage_risk,
        })
    return {
        "report": "Freshness Report",
        "generated_on": str(date.today()),
        "summary": {
            "total_batches": len(batches),
            "assessed_batches": len(scores),
            "fresh_or_good": sum(1 for b in batches if freshness_status(b.freshness_score) in {"Fresh", "Good"}),
            "requires_attention": sum(1 for b in batches if freshness_status(b.freshness_score) in {"Near Spoilage", "Spoiled"}),
            "average_freshness": round(sum(scores) / len(scores), 2) if scores else None,
        },
        "records": records,
    }


def _shelf_life_report(db):
    batches = db.query(Batch).all()
    food_map = get_food_map(db)
    today = date.today()
    records = []
    for batch in batches:
        effective = (
            min(batch.expiry_date, batch.predicted_expiry_date)
            if batch.predicted_expiry_date
            else batch.expiry_date
        )
        days = (effective - today).days
        records.append({
            "batch_id": batch.id,
            "food_name": food_map.get(batch.food_item_id, "Unknown Food"),
            "recorded_expiry": str(batch.expiry_date),
            "predicted_expiry": str(batch.predicted_expiry_date) if batch.predicted_expiry_date else None,
            "effective_expiry": str(effective),
            "days_remaining": days,
            "shelf_life_at_assessment": batch.remaining_shelf_life,
            "current_status": expiry_status(days),
        })
    records.sort(key=lambda r: r["effective_expiry"])
    return {
        "report": "Shelf-Life Report",
        "generated_on": str(today),
        "summary": {
            "total_batches": len(records),
            "fresh": sum(1 for r in records if r["current_status"] == "Fresh"),
            "near_expiry": sum(1 for r in records if r["current_status"] == "Near Expiry"),
            "critical": sum(1 for r in records if r["current_status"] == "Critical"),
            "expired": sum(1 for r in records if r["current_status"] == "Expired"),
        },
        "records": records,
    }


def _inventory_report(db):
    batches = db.query(Batch).all()
    food_map = get_food_map(db)
    today = date.today()

    active_batches = [b for b in batches if (b.quantity or 0) > 0]
    rotation_plan = generate_rotation_plan(active_batches)
    rotation_map = {item["batch_id"]: item for item in rotation_plan}

    records = []
    for batch in batches:
        risk = (
            (batch.freshness_score is not None and batch.freshness_score < 60)
            or (batch.spoilage_risk or "").lower() in {"medium", "high", "very high"}
        )
        effective = (
            min(batch.expiry_date, batch.predicted_expiry_date)
            if batch.predicted_expiry_date
            else batch.expiry_date
        )
        days = (effective - today).days
        rotation = rotation_map.get(batch.id)
        records.append({
            "batch_id": batch.id,
            "food_name": food_map.get(batch.food_item_id, "Unknown Food"),
            "quantity": batch.quantity,
            "availability": "Available" if batch.quantity > 0 else "Out of Stock",
            "freshness_score": batch.freshness_score,
            "spoilage_risk": batch.spoilage_risk,
            "quality_flag": "Requires Attention" if risk else "Normal",
            "effective_expiry": str(effective),
            "days_remaining": days,
            "fefo_priority": rotation.get("rotation_priority") if rotation else None,
            "fefo_action": rotation.get("rotation_action") if rotation else "Out of stock - excluded from FEFO rotation",
        })

    records.sort(key=lambda r: (r["fefo_priority"] is None, r["fefo_priority"] or 99999))
    return {
        "report": "Inventory Quality Report",
        "generated_on": str(today),
        "summary": {
            "total_batches": len(records),
            "available": sum(1 for r in records if r["availability"] == "Available"),
            "out_of_stock": sum(1 for r in records if r["availability"] == "Out of Stock"),
            "quality_risk": sum(1 for r in records if r["quality_flag"] == "Requires Attention"),
            "total_quantity": sum(max(0, int(r["quantity"] or 0)) for r in records),
            "fefo_active_batches": len(rotation_plan),
        },
        "records": records,
    }


def _waste_report(db):
    batches = db.query(Batch).all()
    food_map = get_food_map(db)
    today = date.today()
    records = []
    for batch in batches:
        effective = (
            min(batch.expiry_date, batch.predicted_expiry_date)
            if batch.predicted_expiry_date
            else batch.expiry_date
        )
        days = (effective - today).days
        probability = float(batch.spoilage_probability or 0)
        score = float(batch.freshness_score if batch.freshness_score is not None else 100)
        if days <= 2 or probability >= 60 or score < 40:
            risk, action = "High", "Remove from usable inventory immediately; do not sell or consume."
        elif days <= 7 or probability >= 30 or score < 70:
            risk, action = "Moderate", "Move ahead in FEFO order and consume or sell soon."
        else:
            risk, action = "Low", "Maintain storage and continue routine monitoring."
        records.append({
            "batch_id": batch.id,
            "food_name": food_map.get(batch.food_item_id, "Unknown Food"),
            "quantity": batch.quantity,
            "effective_expiry": str(effective),
            "days_remaining": days,
            "freshness_score": batch.freshness_score,
            "spoilage_probability": batch.spoilage_probability,
            "waste_risk": risk,
            "recommended_action": action,
        })
    records.sort(key=lambda r: r["days_remaining"])
    return {
        "report": "Waste Reduction Report",
        "generated_on": str(today),
        "summary": {
            "total_batches": len(records),
            "high_risk": sum(1 for r in records if r["waste_risk"] == "High"),
            "moderate_risk": sum(1 for r in records if r["waste_risk"] == "Moderate"),
            "low_risk": sum(1 for r in records if r["waste_risk"] == "Low"),
        },
        "records": records,
    }



def _rotation_report(db):
    batches = db.query(Batch).filter(Batch.quantity > 0).all()
    food_map = get_food_map(db)
    plan = generate_rotation_plan(batches)
    records = []
    for item in plan:
        records.append({
            "fefo_priority": item["rotation_priority"],
            "batch_id": item["batch_id"],
            "food_name": food_map.get(item["food_item_id"], "Unknown Food"),
            "quantity": item["quantity"],
            "original_expiry": str(item["original_expiry_date"]),
            "predicted_expiry": str(item["predicted_expiry_date"]) if item["predicted_expiry_date"] else None,
            "effective_expiry": str(item["effective_expiry_date"]),
            "days_until_expiry": item["days_until_expiry"],
            "spoilage_risk": item["spoilage_risk"],
            "rotation_action": item["rotation_action"],
        })
    return {
        "report": "FEFO Inventory Rotation Report",
        "generated_on": str(date.today()),
        "summary": {
            "rotation_method": "First Expired, First Out (FEFO)",
            "active_batches": len(records),
            "expired_active_batches": sum(1 for r in records if r["days_until_expiry"] < 0),
            "immediate_action": sum(1 for r in records if r["days_until_expiry"] <= 2),
            "consume_or_sell_soon": sum(1 for r in records if 3 <= r["days_until_expiry"] <= 5),
        },
        "records": records,
    }

def _storage_report(db):
    readings = db.query(StorageReading).order_by(StorageReading.recorded_at.desc()).all()
    records = []
    for item in readings:
        records.append({
            "reading_id": item.id,
            "batch_id": item.batch_id,
            "food_name": item.food_type,
            "temperature": item.temperature,
            "humidity": item.humidity,
            "air_circulation": item.air_circulation,
            "light_exposure": item.light_exposure,
            "packaging_type": item.packaging_type,
            "storage_duration": item.storage_duration,
            "compliance_score": item.compliance_score,
            "overall_status": item.overall_status,
            "temperature_status": item.temperature_status,
            "humidity_status": item.humidity_status,
            "air_status": item.air_status,
            "light_status": item.light_status,
            "recorded_at": str(item.recorded_at),
        })
    return {
        "report": "Storage Compliance Report",
        "generated_on": str(date.today()),
        "summary": {
            "total_readings": len(records),
            "compliant": sum(1 for r in records if r["compliance_score"] >= 75),
            "needs_attention": sum(1 for r in records if r["compliance_score"] < 75),
            "critical": sum(1 for r in records if r["compliance_score"] < 50),
            "average_compliance": round(sum(r["compliance_score"] for r in records) / len(records), 2) if records else None,
        },
        "records": records,
    }


@router.get("/freshness")
def freshness_report():
    db = SessionLocal()
    try:
        return _freshness_report(db)
    finally:
        db.close()


@router.get("/shelf-life")
def shelf_life_report():
    db = SessionLocal()
    try:
        return _shelf_life_report(db)
    finally:
        db.close()


@router.get("/inventory-quality")
def inventory_quality_report():
    db = SessionLocal()
    try:
        return _inventory_report(db)
    finally:
        db.close()


@router.get("/waste-reduction")
def waste_reduction_report():
    db = SessionLocal()
    try:
        return _waste_report(db)
    finally:
        db.close()


@router.get("/storage-compliance")
def storage_compliance_report():
    db = SessionLocal()
    try:
        return _storage_report(db)
    finally:
        db.close()


@router.get("/fefo-rotation")
def fefo_rotation_report():
    db = SessionLocal()
    try:
        return _rotation_report(db)
    finally:
        db.close()


def _loads(text, fallback):
    try:
        return json.loads(text) if text else fallback
    except Exception:
        return fallback


def _assessment_payload(db, assessment_id: int):
    a = db.query(FreshnessAssessment).filter(FreshnessAssessment.id == assessment_id).first()
    if not a:
        raise HTTPException(status_code=404, detail="Assessment not found")
    batch = db.query(Batch).filter(Batch.id == a.batch_id).first() if a.batch_id else None
    food = db.query(FoodItem).filter(FoodItem.id == batch.food_item_id).first() if batch else None
    today = date.today()
    is_spoiled = str(a.freshness_status or "").strip().lower() == "spoiled"
    days_current = 0 if is_spoiled else (a.predicted_expiry_date - today).days
    current_days_display = (
        "Shelf life exhausted"
        if is_spoiled
        else format_days_remaining(days_current)
    )
    current_status = "Spoiled" if is_spoiled else expiry_status(days_current)
    return {
        "report": "Individual Freshness Analysis Report",
        "report_id": f"FT-{a.recorded_at.strftime('%Y%m%d') if a.recorded_at else date.today().strftime('%Y%m%d')}-{a.id:04d}",
        "generated_on": str(date.today()),
        "assessment_id": a.id,
        "recorded_at": str(a.recorded_at),
        "batch": {
            "batch_id": a.batch_id,
            "food_name": food.name if food else a.food_type,
            "category": food.category if food else None,
            "quantity": batch.quantity if batch else None,
            "recorded_expiry": str(batch.expiry_date) if batch else None,
            "predicted_expiry": str(a.predicted_expiry_date),
            "current_days_remaining": days_current,
            "current_days_display": current_days_display,
            "current_expiry_status": current_status,
        },
        "freshness_analysis": {
            "freshness_score": a.freshness_score,
            "freshness_status": a.freshness_status,
            "shelf_life_at_assessment": a.shelf_life_at_assessment,
            "spoilage_probability": a.spoilage_probability,
            "spoilage_risk": a.spoilage_risk,
            "overall_health_score": a.overall_health_score,
        },
        "storage": {
            "temperature": a.temperature,
            "humidity": a.humidity,
            "air_circulation": a.air_circulation,
            "light_exposure": a.light_exposure,
            "packaging_type": a.packaging_type,
            "storage_duration_days": a.storage_duration,
            "compliance_score": a.storage_compliance_score,
            "overall_status": a.overall_storage_status,
            "details": _loads(a.storage_json, {}),
        },
        "visual_findings": {
            "color": _loads(a.color_analysis_json, {}),
            "texture": _loads(a.texture_analysis_json, {}),
            "mold": _loads(a.mold_analysis_json, {}),
            "bruising": _loads(a.bruising_analysis_json, {}),
            "physical_damage": _loads(a.physical_damage_json, {}),
            "final_assessment": _loads(a.final_assessment_json, {}),
        },
        "shelf_life": _loads(a.shelf_life_json, {}),
        "recommendations": _loads(a.recommendations_json, {}),
        "inventory_rotation": _fefo_entry_for_batch(db, a.batch_id),
        "image_path": a.image_path,
        "image_name": a.image_name,
    }


@router.get("/assessment/{assessment_id}")
def assessment_report(assessment_id: int):
    db = SessionLocal()
    try:
        return _assessment_payload(db, assessment_id)
    finally:
        db.close()


@router.get("/batch/{batch_id}")
def latest_batch_report(batch_id: int):
    db = SessionLocal()
    try:
        assessment = (
            db.query(FreshnessAssessment)
            .filter(FreshnessAssessment.batch_id == batch_id)
            .order_by(FreshnessAssessment.recorded_at.desc())
            .first()
        )
        if assessment:
            return _assessment_payload(db, assessment.id)

        batch = db.query(Batch).filter(Batch.id == batch_id).first()
        if not batch:
            raise HTTPException(status_code=404, detail="Batch not found")
        food = db.query(FoodItem).filter(FoodItem.id == batch.food_item_id).first()
        history = (
            db.query(PredictionHistory)
            .filter(PredictionHistory.batch_id == batch_id)
            .order_by(PredictionHistory.recorded_at.desc())
            .first()
        )
        effective = (
            min(batch.expiry_date, batch.predicted_expiry_date)
            if batch.predicted_expiry_date
            else batch.expiry_date
        )
        days = (effective - date.today()).days
        return {
            "report": "Individual Freshness Analysis Report",
            "generated_on": str(date.today()),
            "assessment_id": None,
            "batch": {
                "batch_id": batch.id,
                "food_name": food.name if food else "Unknown Food",
                "category": food.category if food else None,
                "quantity": batch.quantity,
                "recorded_expiry": str(batch.expiry_date),
                "predicted_expiry": str(batch.predicted_expiry_date) if batch.predicted_expiry_date else None,
                "current_days_remaining": days,
                "current_days_display": format_days_remaining(days),
                "current_expiry_status": expiry_status(days),
            },
            "freshness_analysis": {
                "freshness_score": batch.freshness_score,
                "freshness_status": freshness_status(batch.freshness_score),
                "shelf_life_at_assessment": batch.remaining_shelf_life,
                "spoilage_probability": batch.spoilage_probability,
                "spoilage_risk": batch.spoilage_risk,
            },
            "storage": {
                "temperature": history.temperature if history else None,
                "humidity": history.humidity if history else None,
                "recorded_at": str(history.recorded_at) if history else None,
            },
            "inventory_rotation": _fefo_entry_for_batch(db, batch.id),
            "note": "Run a new Complete Freshness Assessment to include full visual findings, storage compliance, recommendations, PDF image, and Excel detail sheets.",
        }
    finally:
        db.close()


def _flatten_for_excel(report):
    summary = []
    for section in ["batch", "freshness_analysis", "storage", "inventory_rotation"]:
        for k, v in report.get(section, {}).items():
            if isinstance(v, (dict, list)):
                continue
            summary.append((section.replace("_", " ").title(), k.replace("_", " ").title(), v))
    return summary


def format_days_remaining(days):
    """Human-readable expiry wording for reports; never expose raw negative days."""
    if days is None:
        return "-"
    try:
        days = int(days)
    except (TypeError, ValueError):
        return str(days)
    if days < 0:
        return f"Expired {abs(days)} day(s) ago"
    if days == 0:
        return "Expires today"
    return f"{days} day(s) left"


def _assessment_pdf_bytes(report):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT, TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, KeepTogether, PageBreak
    )

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4, rightMargin=14*mm, leftMargin=14*mm,
        topMargin=13*mm, bottomMargin=13*mm,
        title="FoodTrace Individual Freshness Analysis Report",
    )
    styles = getSampleStyleSheet()
    brand = colors.HexColor("#0A4F3F")
    green = colors.HexColor("#14945F")
    navy = colors.HexColor("#102A2E")
    blue = colors.HexColor("#5B8FF9")
    amber = colors.HexColor("#F3B64A")
    red = colors.HexColor("#E25555")
    pale = colors.HexColor("#EEF4F8")
    line = colors.HexColor("#D7E3DF")

    title = ParagraphStyle("FTTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=20, leading=24, textColor=navy, alignment=TA_LEFT, spaceAfter=5)
    h2 = ParagraphStyle("FTH2", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=12, leading=15, textColor=brand, spaceBefore=7, spaceAfter=7)
    body = ParagraphStyle("FTBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=8.5, leading=12, textColor=colors.HexColor("#40545B"))
    small = ParagraphStyle("FTSmall", parent=body, fontSize=7.4, leading=10)
    value = ParagraphStyle("FTValue", parent=body, fontName="Helvetica-Bold", fontSize=13, leading=15, textColor=navy)

    story = []
    header = Table([
        [Paragraph("<b><font color='#0A4F3F'>Food</font><font color='#14945F'>Trace</font></b><br/><font size='7'>TRACE TODAY. FRESHER TOMORROW.</font>", body),
         Paragraph("AI-POWERED FOOD FRESHNESS MONITORING<br/><font size='7'>Individual quality intelligence report</font>", small)]
    ], colWidths=[92*mm, 85*mm])
    header.setStyle(TableStyle([("VALIGN",(0,0),(-1,-1),"TOP"),("ALIGN",(1,0),(1,0),"RIGHT"),("LINEBELOW",(0,0),(-1,-1),1,brand),("BOTTOMPADDING",(0,0),(-1,-1),6)]))
    story += [header, Spacer(1,5), Paragraph("Individual Freshness Analysis Report", title)]
    story.append(Paragraph(f"Report ID: {report.get('report_id','FoodTrace')} &nbsp;&nbsp; | &nbsp;&nbsp; Generated: {report.get('generated_on')}", small))
    story.append(Spacer(1,8))

    batch = report.get("batch", {})
    fa = report.get("freshness_analysis", {})
    storage = report.get("storage", {})
    recs = report.get("recommendations", {})
    rotation = report.get("inventory_rotation") or {}

    img_path = report.get("image_path")
    if img_path:
        try:
            # Assessments normally save a path relative to the backend working directory.
            # Resolve it safely as-is first, then relative to the backend project root.
            candidates = [img_path]
            backend_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
            if not os.path.isabs(img_path):
                candidates.append(os.path.join(backend_root, img_path))
            existing = next((candidate for candidate in candidates if os.path.exists(candidate)), None)

            if existing:
                safe_img_path = existing
                if not existing.lower().endswith((".png", ".jpg", ".jpeg")):
                    from PIL import Image as PILImage
                    converted_path = f"{existing}.report.png"
                    with PILImage.open(existing) as source:
                        source.convert("RGB").save(converted_path, "PNG")
                    safe_img_path = converted_path

                photo = Image(safe_img_path, width=38*mm, height=28*mm)
                photo.hAlign = "LEFT"
                photo_table = Table([[photo, Paragraph(
                    f"<b>Uploaded food image</b><br/>{batch.get('food_name','-')}"
                    f"<br/>Assessment #{report.get('assessment_id','-')}", body)]],
                    colWidths=[43*mm,134*mm])
                photo_table.setStyle(TableStyle([("VALIGN",(0,0),(-1,-1),"MIDDLE"),("BOX",(0,0),(-1,-1),.5,line),("BACKGROUND",(0,0),(-1,-1),colors.white),("LEFTPADDING",(0,0),(-1,-1),6),("RIGHTPADDING",(0,0),(-1,-1),6),("TOPPADDING",(0,0),(-1,-1),5),("BOTTOMPADDING",(0,0),(-1,-1),5)]))
                story.append(photo_table)
                story.append(Spacer(1,6))
            else:
                story.append(Paragraph(
                    f"<b>Uploaded food image:</b> {report.get('image_name') or 'Image file not available in this environment'}",
                    body,
                ))
                story.append(Spacer(1,4))
        except Exception:
            story.append(Paragraph(
                f"<b>Uploaded food image:</b> {report.get('image_name') or 'Image could not be embedded'}",
                body,
            ))
            story.append(Spacer(1,4))

    info_rows = [
        ["Food", batch.get("food_name") or "-", "Batch", f"#{batch.get('batch_id')}" if batch.get("batch_id") else "Unlinked"],
        ["Category", batch.get("category") or "-", "Quantity", f"{batch.get('quantity')} kg" if batch.get("quantity") is not None else "-"],
        ["Recorded Expiry", batch.get("recorded_expiry") or "-", "Predicted Expiry", batch.get("predicted_expiry") or "-"],
        ["Current Status", batch.get("current_expiry_status") or "-", "Expiry Timing", batch.get("current_days_display") or format_days_remaining(batch.get("current_days_remaining"))],
    ]
    story.append(Paragraph("1. Batch Information", h2))
    t = Table(info_rows, colWidths=[29*mm,58*mm,31*mm,59*mm], hAlign="LEFT")
    t.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,-1),pale),("GRID",(0,0),(-1,-1),.35,line),("FONTNAME",(0,0),(-1,-1),"Helvetica"),("FONTNAME",(0,0),(0,-1),"Helvetica-Bold"),("FONTNAME",(2,0),(2,-1),"Helvetica-Bold"),("FONTSIZE",(0,0),(-1,-1),8),("TEXTCOLOR",(0,0),(-1,-1),navy),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("TOPPADDING",(0,0),(-1,-1),6),("BOTTOMPADDING",(0,0),(-1,-1),6)]))
    story.append(t)

    story.append(Paragraph("2. Executive Summary", h2))
    metrics = [
        ("Freshness Score", f"{fa.get('freshness_score','-')}%", green),
        ("Freshness Status", str(fa.get('freshness_status','-')), green),
        ("Shelf Life at Assessment", f"{fa.get('shelf_life_at_assessment','-')} days", blue),
        ("Spoilage Risk", str(fa.get('spoilage_risk','-')), amber),
        ("Spoilage Probability", f"{fa.get('spoilage_probability','-')}%", red if float(fa.get('spoilage_probability') or 0)>=50 else amber),
        ("Overall Food Health", f"{fa.get('overall_health_score','-')}%", green),
    ]
    metric_cells = []
    for label, val, accent in metrics:
        metric_cells.append(Table([[Paragraph(label, small)],[Paragraph(str(val), value)]], colWidths=[56*mm], style=TableStyle([("BACKGROUND",(0,0),(-1,-1),colors.white),("BOX",(0,0),(-1,-1),1,accent),("LEFTPADDING",(0,0),(-1,-1),8),("TOPPADDING",(0,0),(-1,-1),6),("BOTTOMPADDING",(0,0),(-1,-1),6)])))
    story.append(Table([metric_cells[:3],metric_cells[3:]], colWidths=[59*mm]*3, hAlign="LEFT", style=TableStyle([("VALIGN",(0,0),(-1,-1),"TOP"),("LEFTPADDING",(0,0),(-1,-1),0),("RIGHTPADDING",(0,0),(-1,-1),3),("TOPPADDING",(0,0),(-1,-1),2),("BOTTOMPADDING",(0,0),(-1,-1),2)])))

    story.append(Paragraph("3. Storage Condition Compliance", h2))
    details = storage.get("details") or {}
    srows = [["Condition","Actual","Recommended","Status"]]
    for key,label in [("temperature","Temperature"),("humidity","Humidity"),("air_circulation","Air Circulation"),("light_exposure","Light Exposure")]:
        d = details.get(key, {}) if isinstance(details, dict) else {}
        actual = d.get("current")
        if key == "temperature" and actual is None: actual = storage.get("temperature")
        if key == "humidity" and actual is None: actual = storage.get("humidity")
        if key == "air_circulation" and actual is None: actual = storage.get("air_circulation")
        if key == "light_exposure" and actual is None: actual = storage.get("light_exposure")
        rec = d.get("recommended", "-")
        if isinstance(rec, list): rec = ", ".join(rec)
        srows.append([label, str(actual if actual is not None else "-"), str(rec), str(d.get("status","-"))])
    srows.append(["Packaging", str(storage.get("packaging_type") or "-"), "Food-appropriate protected packaging", "Recorded"])
    srows.append(["Storage Duration", f"{storage.get('storage_duration_days',0)} days", "Monitor continuously", "Recorded"])
    st = Table(srows, colWidths=[42*mm,38*mm,61*mm,36*mm])
    st.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),brand),("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("GRID",(0,0),(-1,-1),.35,line),("FONTSIZE",(0,0),(-1,-1),7.5),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,pale]),("TOPPADDING",(0,0),(-1,-1),5),("BOTTOMPADDING",(0,0),(-1,-1),5)]))
    story.append(st)
    story.append(Paragraph(f"Overall storage compliance: <b>{storage.get('compliance_score','-')}%</b> - {storage.get('overall_status','-')}", body))

    story.append(Paragraph("4. Visual Freshness Findings", h2))
    vf = report.get("visual_findings", {})
    def compact(d):
        if not isinstance(d, dict): return str(d)
        parts=[]
        for k,v in d.items():
            if isinstance(v,(str,int,float,bool)) and k not in {"confidence"}:
                parts.append(f"{k.replace('_',' ').title()}: {v}")
            if len(parts)>=3: break
        return "; ".join(parts) or "Analysis recorded"
    vrows=[["Finding","Observation"],
           ["Color",compact(vf.get("color",{}))],
           ["Texture",compact(vf.get("texture",{}))],
           ["Mold",compact(vf.get("mold",{}))],
           ["Bruising",compact(vf.get("bruising",{}))],
           ["Physical Damage",compact(vf.get("physical_damage",{}))]]
    vt=Table(vrows,colWidths=[43*mm,134*mm])
    vt.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#DDEBE6")),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("GRID",(0,0),(-1,-1),.35,line),("FONTSIZE",(0,0),(-1,-1),7.2),("VALIGN",(0,0),(-1,-1),"TOP"),("TOPPADDING",(0,0),(-1,-1),5),("BOTTOMPADDING",(0,0),(-1,-1),5)]))
    story.append(vt)

    story.append(Paragraph("5. Shelf-Life Trend", h2))
    try:
        from reportlab.graphics.shapes import Drawing, Line, Circle, String, PolyLine
        actual_days = max(0, int(fa.get("shelf_life_at_assessment") or 0))
        drawing = Drawing(500, 125)
        x0, y0, w, h = 34, 24, 430, 78
        drawing.add(Line(x0, y0, x0+w, y0, strokeColor=colors.HexColor("#AFC4BC"), strokeWidth=1))
        drawing.add(Line(x0, y0, x0, y0+h, strokeColor=colors.HexColor("#AFC4BC"), strokeWidth=1))

        if actual_days <= 0:
            drawing.add(Circle(x0, y0, 2.8, fillColor=green, strokeColor=colors.white, strokeWidth=.7))
            drawing.add(String(2, y0-2, "0d", fontName="Helvetica", fontSize=7, fillColor=colors.HexColor("#60747B")))
            drawing.add(String(x0,8,"Assessment day",fontName="Helvetica",fontSize=7,fillColor=colors.HexColor("#60747B")))
            drawing.add(String(x0+w-90,8,"Shelf life exhausted",fontName="Helvetica",fontSize=7,fillColor=colors.HexColor("#60747B")))
        else:
            pts = []
            steps = min(actual_days, 14)
            for i in range(steps + 1):
                day = (actual_days * i / steps) if steps else 0
                remaining = max(actual_days - day, 0)
                x = x0 + (i / steps) * w if steps else x0
                y = y0 + (remaining / actual_days) * h if actual_days else y0
                pts.extend([x,y])
                drawing.add(Circle(x,y,2.4,fillColor=green,strokeColor=colors.white,strokeWidth=.7))
            drawing.add(PolyLine(pts,strokeColor=green,strokeWidth=2,fillColor=None))
            drawing.add(String(x0,8,"Assessment day",fontName="Helvetica",fontSize=7,fillColor=colors.HexColor("#60747B")))
            drawing.add(String(x0+w-90,8,"Predicted expiry",fontName="Helvetica",fontSize=7,fillColor=colors.HexColor("#60747B")))
            drawing.add(String(2,y0+h-4,f"{actual_days}d",fontName="Helvetica",fontSize=7,fillColor=colors.HexColor("#60747B")))
            drawing.add(String(12,y0-2,"0d",fontName="Helvetica",fontSize=7,fillColor=colors.HexColor("#60747B")))

        story.append(drawing)
        shelf_text = "Shelf life exhausted" if actual_days <= 0 else f"{actual_days} day(s)"
        story.append(Paragraph(
            f"Shelf life at assessment: <b>{shelf_text}</b>. Predicted expiry: <b>{batch.get('predicted_expiry') or '-'}</b>. "
            f"Current expiry status: <b>{batch.get('current_expiry_status') or '-'}</b>.", body))
    except Exception:
        story.append(Paragraph("Shelf-life trend data is available in the digital report.", body))

    rec_heading = Paragraph("6. Recommendations", h2)
    rec_sections = [
        ("Storage", recs.get("storage_recommendations",[])),
        ("Consumption", recs.get("consumption_recommendations",[])),
        ("Waste Reduction", recs.get("waste_reduction_recommendations",[])),
        ("Quality Improvement", recs.get("quality_improvement_recommendations",[])),
    ]
    rrows=[]
    for label, items in rec_sections:
        if isinstance(items,list): text="; ".join(str(x) for x in items)
        else: text=str(items or "-")
        rrows.append([Paragraph(f"<b>{label}</b>",body),Paragraph(text,body)])
    rrows.append([Paragraph("<b>Overall Action</b>",body),Paragraph(str(recs.get("overall_action","Continue monitoring")),body)])
    rt=Table(rrows,colWidths=[40*mm,137*mm])
    rt.setStyle(TableStyle([("GRID",(0,0),(-1,-1),.35,line),("BACKGROUND",(0,0),(0,-1),pale),("VALIGN",(0,0),(-1,-1),"TOP"),("TOPPADDING",(0,0),(-1,-1),6),("BOTTOMPADDING",(0,0),(-1,-1),6)]))
    story.append(KeepTogether([rec_heading, rt]))

    story.append(Paragraph("7. FEFO Inventory Rotation", h2))
    is_spoiled = str(fa.get("freshness_status") or "").strip().lower() == "spoiled" or str(batch.get("current_expiry_status") or "").strip().lower() == "spoiled"
    if is_spoiled:
        expiry_timing = "Shelf life exhausted"
        rotation_action = "Remove from usable inventory; do not rotate for consumption."
    else:
        expiry_timing = format_days_remaining(rotation.get("days_until_expiry") if rotation.get("days_until_expiry") is not None else batch.get("current_days_remaining"))
        rotation_action = rotation.get("rotation_action") or "Continue FEFO rotation based on effective expiry."

    fefo_rows = [
        ["Rotation Method", rotation.get("rotation_method") or "FEFO - First Expired, First Out"],
        ["Priority Rank", rotation.get("priority_rank") if rotation.get("priority_rank") is not None else "-"],
        ["Effective Expiry", rotation.get("effective_expiry_date") or batch.get("predicted_expiry") or batch.get("recorded_expiry") or "-"],
        ["Expiry Timing", expiry_timing],
        ["Rotation Action", rotation_action],
    ]
    fefo_table = Table(fefo_rows, colWidths=[48*mm,129*mm])
    fefo_table.setStyle(TableStyle([
        ("GRID",(0,0),(-1,-1),.35,line),
        ("BACKGROUND",(0,0),(0,-1),pale),
        ("FONTNAME",(0,0),(0,-1),"Helvetica-Bold"),
        ("TEXTCOLOR",(0,0),(-1,-1),navy),
        ("FONTSIZE",(0,0),(-1,-1),7.8),
        ("VALIGN",(0,0),(-1,-1),"TOP"),
        ("TOPPADDING",(0,0),(-1,-1),6),
        ("BOTTOMPADDING",(0,0),(-1,-1),6),
    ]))
    story.append(fefo_table)

    conclusion_heading = Paragraph("8. Final Conclusion", h2)
    status = batch.get("current_expiry_status") or "Monitor"
    conclusion = (
        f"Freshness status is {fa.get('freshness_status','-')} with a score of {fa.get('freshness_score','-')}%. "
        f"Spoilage risk is {fa.get('spoilage_risk','-')} ({fa.get('spoilage_probability','-')}%). "
        f"Current expiry status is {status}. Storage compliance is {storage.get('compliance_score','-')}%. "
        f"Follow the recommended storage, consumption, rotation and waste-reduction actions shown above."
    )
    concl = Table([[Paragraph(conclusion,body)]], colWidths=[177*mm])
    concl.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,-1),colors.HexColor("#E7F4EE")),("BOX",(0,0),(-1,-1),1,green),("LEFTPADDING",(0,0),(-1,-1),9),("RIGHTPADDING",(0,0),(-1,-1),9),("TOPPADDING",(0,0),(-1,-1),8),("BOTTOMPADDING",(0,0),(-1,-1),8)]))
    story.append(KeepTogether([conclusion_heading, concl]))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setStrokeColor(brand); canvas.setLineWidth(.5); canvas.line(14*mm,10*mm,196*mm,10*mm)
        canvas.setFillColor(brand); canvas.setFont("Helvetica-Bold",8); canvas.drawString(14*mm,6*mm,"FoodTrace")
        canvas.setFillColor(colors.HexColor("#60747B")); canvas.setFont("Helvetica",7); canvas.drawRightString(196*mm,6*mm,f"Page {doc.page} | Generated {report.get('generated_on')}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    buf.seek(0)
    return buf.getvalue()


@router.get("/assessment/{assessment_id}/pdf")
def assessment_pdf(assessment_id: int):
    db = SessionLocal()
    try:
        report = _assessment_payload(db, assessment_id)
        pdf = _assessment_pdf_bytes(report)
        filename = f"FoodTrace_Assessment_{assessment_id}.pdf"
        return StreamingResponse(io.BytesIO(pdf), media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="{filename}"'})
    finally:
        db.close()


@router.get("/assessment/{assessment_id}/excel")
def assessment_excel(assessment_id: int):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

    db = SessionLocal()
    try:
        report = _assessment_payload(db, assessment_id)
    finally:
        db.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    brand = "0A4F3F"; green = "14945F"; pale = "EEF4F8"; white = "FFFFFF"; line = "D7E3DF"
    ws["A1"] = "FoodTrace - Individual Freshness Analysis Report"
    ws["A1"].font = Font(size=18, bold=True, color=brand)
    ws.merge_cells("A1:C1")
    ws["A2"] = f"Report ID: {report.get('report_id')}"
    ws["A3"] = f"Generated: {report.get('generated_on')}"
    ws.append([])
    ws.append(["Section", "Field", "Value"])
    for cell in ws[5]:
        cell.fill = PatternFill("solid", fgColor=brand); cell.font = Font(bold=True,color=white)
    for row in _flatten_for_excel(report):
        ws.append(list(row))
    for col,width in {"A":24,"B":30,"C":42}.items(): ws.column_dimensions[col].width=width
    for row in ws.iter_rows(min_row=6):
        for c in row:
            c.alignment=Alignment(vertical="top",wrap_text=True)

    storage = wb.create_sheet("Storage Compliance")
    storage.append(["Condition","Actual","Recommended","Status"])
    for c in storage[1]: c.fill=PatternFill("solid",fgColor=brand); c.font=Font(bold=True,color=white)
    details = report.get("storage",{}).get("details",{})
    for key,label in [("temperature","Temperature"),("humidity","Humidity"),("air_circulation","Air Circulation"),("light_exposure","Light Exposure")]:
        d=details.get(key,{}) if isinstance(details,dict) else {}
        rec=d.get("recommended","-")
        if isinstance(rec,list): rec=", ".join(rec)
        storage.append([label,d.get("current","-"),rec,d.get("status","-")])
    storage.append(["Compliance Score",report.get("storage",{}).get("compliance_score"),"75% or higher preferred",report.get("storage",{}).get("overall_status")])
    for col,width in {"A":24,"B":22,"C":38,"D":24}.items(): storage.column_dimensions[col].width=width

    rec = wb.create_sheet("Recommendations")
    rec.append(["Recommendation Type","Action"])
    for c in rec[1]: c.fill=PatternFill("solid",fgColor=brand); c.font=Font(bold=True,color=white)
    r=report.get("recommendations",{})
    for label,key in [("Storage","storage_recommendations"),("Consumption","consumption_recommendations"),("Waste Reduction","waste_reduction_recommendations"),("Quality Improvement","quality_improvement_recommendations")]:
        items=r.get(key,[])
        if not isinstance(items,list): items=[items]
        for item in items or ["-"]: rec.append([label,item])
    rec.append(["Overall Action",r.get("overall_action","-")])
    rec.column_dimensions["A"].width=28; rec.column_dimensions["B"].width=80

    fefo = wb.create_sheet("FEFO Rotation")
    fefo.append(["Field", "Value"])
    for c in fefo[1]: c.fill=PatternFill("solid",fgColor=brand); c.font=Font(bold=True,color=white)
    rotation = report.get("inventory_rotation") or {}
    for key in ["rotation_method", "priority_rank", "effective_expiry_date", "days_until_expiry", "rotation_action", "availability"]:
        fefo.append([key.replace("_", " ").title(), rotation.get(key, "-")])
    fefo.column_dimensions["A"].width=30; fefo.column_dimensions["B"].width=75

    visual=wb.create_sheet("Visual Findings")
    visual.append(["Finding","Details"])
    for c in visual[1]: c.fill=PatternFill("solid",fgColor=brand); c.font=Font(bold=True,color=white)
    for key,val in report.get("visual_findings",{}).items(): visual.append([key.replace("_"," ").title(), json.dumps(val, indent=2)])
    visual.column_dimensions["A"].width=28; visual.column_dimensions["B"].width=90
    for row in visual.iter_rows(min_row=2): row[1].alignment=Alignment(wrap_text=True,vertical="top")

    out=io.BytesIO(); wb.save(out); out.seek(0)
    return StreamingResponse(out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition":f'attachment; filename="FoodTrace_Assessment_{assessment_id}.xlsx"'})


def _management_excel(report):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb=Workbook(); ws=wb.active; ws.title="Report"
    ws["A1"]=report["report"]; ws["A1"].font=Font(size=18,bold=True,color="0A4F3F")
    ws["A2"]=f"Generated: {report['generated_on']}"
    ws.append([]); ws.append(["Summary Metric","Value"])
    for c in ws[4]: c.fill=PatternFill("solid",fgColor="0A4F3F"); c.font=Font(bold=True,color="FFFFFF")
    for k,v in report["summary"].items(): ws.append([k.replace("_"," ").title(),v])
    ws.append([])
    records=report.get("records",[])
    if records:
        headers=list(records[0].keys()); ws.append(headers)
        for c in ws[ws.max_row]: c.fill=PatternFill("solid",fgColor="14945F"); c.font=Font(bold=True,color="FFFFFF")
        for item in records: ws.append([item.get(h) for h in headers])
    for col in ws.columns:
        letter=col[0].column_letter; ws.column_dimensions[letter].width=min(42,max(12,max(len(str(c.value or "")) for c in col)+2))
        for c in col: c.alignment=Alignment(wrap_text=True,vertical="top")
    out=io.BytesIO(); wb.save(out); out.seek(0); return out


REPORT_BUILDERS={
    "freshness": _freshness_report,
    "shelf-life": _shelf_life_report,
    "inventory-quality": _inventory_report,
    "waste-reduction": _waste_report,
    "storage-compliance": _storage_report,
    "fefo-rotation": _rotation_report,
}


@router.get("/{report_type}/excel")
def management_excel(report_type: str):
    builder=REPORT_BUILDERS.get(report_type)
    if not builder: raise HTTPException(status_code=404,detail="Unknown report type")
    db=SessionLocal()
    try: report=builder(db)
    finally: db.close()
    out=_management_excel(report)
    return StreamingResponse(out,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":f'attachment; filename="FoodTrace_{report_type}.xlsx"'})


def _management_pdf(report):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    buf=io.BytesIO(); pagesize=landscape(A4)
    doc=SimpleDocTemplate(buf,pagesize=pagesize,rightMargin=12*mm,leftMargin=12*mm,topMargin=12*mm,bottomMargin=12*mm,title=report["report"])
    styles=getSampleStyleSheet(); navy=colors.HexColor("#102A2E"); brand=colors.HexColor("#0A4F3F"); green=colors.HexColor("#14945F"); pale=colors.HexColor("#EEF4F8"); line=colors.HexColor("#D7E3DF")
    title=ParagraphStyle("title",parent=styles["Title"],fontName="Helvetica-Bold",fontSize=19,textColor=navy,alignment=0)
    body=ParagraphStyle("body",parent=styles["BodyText"],fontSize=7.2,leading=9,textColor=colors.HexColor("#40545B"))
    table_header=ParagraphStyle("table_header",parent=body,fontName="Helvetica-Bold",fontSize=7.1,leading=8.5,textColor=colors.white)
    story=[Paragraph("FoodTrace",ParagraphStyle("brand",parent=title,fontSize=16,textColor=brand)),Paragraph(report["report"],title),Paragraph(f"Generated: {report['generated_on']}",body),Spacer(1,7)]
    summary=[[k.replace("_"," ").title(),str(v)] for k,v in report.get("summary",{}).items()]
    if summary:
        st=Table(summary,colWidths=[48*mm,36*mm],hAlign="LEFT")
        st.setStyle(TableStyle([("BACKGROUND",(0,0),(0,-1),pale),("GRID",(0,0),(-1,-1),.35,line),("FONTNAME",(0,0),(0,-1),"Helvetica-Bold"),("FONTSIZE",(0,0),(-1,-1),7.5),("TOPPADDING",(0,0),(-1,-1),4),("BOTTOMPADDING",(0,0),(-1,-1),4)])); story += [st,Spacer(1,8)]
    records=report.get("records",[])
    if records:
        headers=list(records[0].keys())
        data=[[Paragraph(h.replace("_"," ").title(),table_header) for h in headers]]
        for item in records[:120]: data.append([Paragraph(str(item.get(h,"-")),body) for h in headers])
        avail=270*mm; widths=[avail/len(headers)]*len(headers)
        t=Table(data,colWidths=widths,repeatRows=1)
        t.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),brand),("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("GRID",(0,0),(-1,-1),.25,line),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,pale]),("VALIGN",(0,0),(-1,-1),"TOP"),("TOPPADDING",(0,0),(-1,-1),3),("BOTTOMPADDING",(0,0),(-1,-1),3)])); story.append(t)
    def footer(canvas,doc):
        canvas.saveState(); canvas.setStrokeColor(brand); canvas.line(12*mm,8*mm,285*mm,8*mm); canvas.setFont("Helvetica",7); canvas.setFillColor(brand); canvas.drawString(12*mm,4*mm,"FoodTrace - AI Food Freshness Monitoring"); canvas.drawRightString(285*mm,4*mm,f"Page {doc.page}"); canvas.restoreState()
    doc.build(story,onFirstPage=footer,onLaterPages=footer); buf.seek(0); return buf


@router.get("/{report_type}/pdf")
def management_pdf(report_type: str):
    builder=REPORT_BUILDERS.get(report_type)
    if not builder: raise HTTPException(status_code=404,detail="Unknown report type")
    db=SessionLocal()
    try: report=builder(db)
    finally: db.close()
    out=_management_pdf(report)
    return StreamingResponse(out,media_type="application/pdf",headers={"Content-Disposition":f'attachment; filename="FoodTrace_{report_type}.pdf"'})
