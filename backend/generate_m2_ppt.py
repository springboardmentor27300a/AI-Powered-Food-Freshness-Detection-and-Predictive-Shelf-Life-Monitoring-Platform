import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_PARAGRAPH_ALIGNMENT

METRICS_PATH = 'app/ml/saved_models/metrics.json'
OUT = 'FreshGuard_M2_Review_Presentation.pptx'

with open(METRICS_PATH, 'r', encoding='utf-8') as f:
    metrics = json.load(f)

prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)

# Helper to add title + content slide
def add_title_slide(title, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    title_box = slide.shapes.title
    if title_box:
        title_box.text = title
    else:
        tx = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.3), Inches(1))
        p = tx.text_frame.add_paragraph()
        p.text = title
        p.font.size = Pt(32)
    if subtitle:
        tx = slide.shapes.add_textbox(Inches(0.5), Inches(1.0), Inches(12.3), Inches(1))
        p = tx.text_frame.add_paragraph()
        p.text = subtitle
        p.font.size = Pt(20)
    return slide

# 1. Title
slide = add_title_slide('FreshGuard – Milestone 2 Review', 'Image Analysis & Freshness Assessment')

# 2. M2 Objectives
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = 'M2 Objectives'
body = slide.shapes.placeholders[1].text_frame
body.clear()
for item in ['Image analysis', 'Freshness classification', 'Food quality scoring', 'Spoilage detection', 'Freshness assessment/report']:
    p = body.add_paragraph()
    p.text = f'- {item}'
    p.level = 0

# 3. M2 Architecture
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = 'M2 Architecture'
body = slide.shapes.placeholders[1].text_frame
body.clear()
arch = 'Image → MobileNetV2 + OpenCV → Assessment Engine → Freshness Result'
p = body.add_paragraph()
p.text = arch
p.level = 0

# 4. Technology & Dataset
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = 'Technology & Dataset'
body = slide.shapes.placeholders[1].text_frame
body.clear()
tech_items = [
    'MobileNetV2', 'OpenCV', 'PyTorch', 'FastAPI', 'React', 'PostgreSQL',
    'Fruits Fresh and Rotten dataset', f"{metrics['dataset_summary']['total_images']} images", f"{metrics['dataset_summary']['total_classes']} classes"
]
for t in tech_items:
    p = body.add_paragraph()
    p.text = f'- {t}'

# 5. Results
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = 'Results'
body = slide.shapes.placeholders[1].text_frame
body.clear()
# Use metrics from metrics.json
ta = metrics.get('test_accuracy')
pp = metrics.get('test_precision_macro')
rp = metrics.get('test_recall_macro')
f1 = metrics.get('test_f1_macro')
if ta is not None:
    p = body.add_paragraph(); p.text = f"Test Accuracy: {ta * 100:.2f}%"; p.level = 0
if pp is not None:
    p = body.add_paragraph(); p.text = f"Macro Precision: {pp * 100:.2f}%"; p.level = 0
if rp is not None:
    p = body.add_paragraph(); p.text = f"Macro Recall: {rp * 100:.2f}%"; p.level = 0
if f1 is not None:
    p = body.add_paragraph(); p.text = f"Macro F1: {f1 * 100:.2f}%"; p.level = 0
# Note about screenshots
p = body.add_paragraph()
if False:
    p.text = 'Included screenshots from application.'
else:
    p.text = 'No UI screenshots available in the workspace; dataset/example images used where applicable.'

# 6. Live Demo
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = 'Live Demo Flow'
body = slide.shapes.placeholders[1].text_frame
body.clear()
flow = ['Upload image', 'Analyze', 'Freshness classification', 'Confidence', 'Quality score', 'Spoilage probability', 'Analysis history']
for s in flow:
    p = body.add_paragraph(); p.text = f'- {s}'; p.level = 0

# Save
prs.save(OUT)
print('Saved', OUT)
