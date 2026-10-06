"""
/api/storage — Milestone 3 storage condition monitoring.
"""
import uuid

from fastapi import APIRouter, Depends, Query
from sqlalchemy.orm import Session

from app.dependencies.auth import get_current_user
from app.dependencies.db import get_db
from app.dependencies.roles import require_role
from app.models.user import User, UserRole
from app.models.storage import StorageReading
from app.schemas.storage import (StorageReadingCreate, StorageReadingListResponse,
                                  StorageReadingOut, StorageTrendPoint, StorageTrendResponse)
from app.services import storage_service
from io import BytesIO
import csv
from fastapi.responses import Response


from openpyxl import Workbook
from openpyxl.styles import Font



router = APIRouter(prefix="/api/storage", tags=["storage"])

LOGGER_ROLES = (UserRole.WAREHOUSE_OPERATOR, UserRole.ADMINISTRATOR)


@router.post("/readings", response_model=StorageReadingOut, status_code=201)
def add_reading(
    data: StorageReadingCreate,
    current_user: User = Depends(require_role(*LOGGER_ROLES)),
    db: Session = Depends(get_db),
):
    return storage_service.create_reading(db, data, current_user.id)


@router.get("/readings", response_model=StorageReadingListResponse)
def list_readings(
    batch_id: uuid.UUID | None = None,
    storage_location: str | None = None,
    limit: int = Query(100, ge=1, le=500),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    items, total = storage_service.list_readings(db, batch_id, storage_location, limit)
    return StorageReadingListResponse(items=items, total=total)


@router.get("/trends", response_model=StorageTrendResponse)
def get_trends(
    storage_location: str,
    days: int = Query(7, ge=1, le=90),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    readings = storage_service.trend_for_location(db, storage_location, days)
    points = [StorageTrendPoint(recorded_at=r.recorded_at, temperature_c=r.temperature_c,
                                 humidity_pct=r.humidity_pct) for r in readings]
    return StorageTrendResponse(storage_location=storage_location, points=points)
@router.get("/alerts")
def get_storage_alerts(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    return storage_service.storage_alerts(db)

@router.get("/compliance-report")
def get_storage_compliance_report(
    limit: int = Query(500, ge=1, le=1000),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """
    Storage compliance report based on actual recorded readings.
    """

    readings = (
        db.query(StorageReading)
        .order_by(StorageReading.recorded_at.desc())
        .limit(limit)
        .all()
    )

    items = []

    for reading in readings:
        items.append(
            {
                "id": str(reading.id),
                "batch_id": (
                    str(reading.batch_id)
                    if reading.batch_id
                    else None
                ),
                "storage_location": reading.storage_location,
                "temperature_c": reading.temperature_c,
                "humidity_pct": reading.humidity_pct,
                "air_circulation": (
                    reading.air_circulation.value
                    if reading.air_circulation
                    else None
                ),
                "light_exposure": (
                    reading.light_exposure.value
                    if reading.light_exposure
                    else None
                ),
                "is_compliant": reading.is_compliant,
                "compliance_notes": reading.compliance_notes,
                "recorded_at": reading.recorded_at,
            }
        )

    compliant_count = sum(
        1 for item in items if item["is_compliant"]
    )

    non_compliant_count = sum(
        1 for item in items if not item["is_compliant"]
    )

    compliance_percentage = (
        round((compliant_count / len(items)) * 100, 2)
        if items
        else 0
    )

    return {
        "total_readings": len(items),
        "compliant_readings": compliant_count,
        "non_compliant_readings": non_compliant_count,
        "compliance_percentage": compliance_percentage,
        "items": items,

    }
    
@router.get("/compliance-report/csv")
def download_storage_compliance_csv(
    limit: int = Query(500, ge=1, le=1000),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    data = get_storage_compliance_report(limit, current_user, db)

    csv_buffer = io.StringIO()
    writer = csv.writer(csv_buffer)

    writer.writerow([
        "Storage Location",
        "Temperature (C)",
        "Humidity (%)",
        "Air Circulation",
        "Light Exposure",
        "Compliant",
        "Compliance Notes",
        "Recorded At",
    ])

    for item in data["items"]:
        writer.writerow([
            item["storage_location"],
            item["temperature_c"],
            item["humidity_pct"],
            item["air_circulation"],
            item["light_exposure"],
            item["is_compliant"],
            item["compliance_notes"],
            item["recorded_at"],
        ])

    return Response(
        content=csv_buffer.getvalue(),
        media_type="text/csv",
        headers={
            "Content-Disposition":
                'attachment; filename="FoodCare_Storage_Compliance_Report.csv"'
        },
    )


@router.get("/compliance-report/excel")
def download_storage_compliance_excel(
    limit: int = Query(500, ge=1, le=1000),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    data = get_storage_compliance_report(limit, current_user, db)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Storage Compliance"

    headers = [
        "Storage Location",
        "Temperature (C)",
        "Humidity (%)",
        "Air Circulation",
        "Light Exposure",
        "Compliant",
        "Compliance Notes",
        "Recorded At",
    ]

    for column, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=column, value=header)
        cell.font = Font(bold=True)

    for row_number, item in enumerate(data["items"], start=2):
        values = [
            item["storage_location"],
            item["temperature_c"],
            item["humidity_pct"],
            item["air_circulation"],
            item["light_exposure"],
            item["is_compliant"],
            item["compliance_notes"],
            str(item["recorded_at"]),
        ]

        for column, value in enumerate(values, start=1):
            worksheet.cell(
                row=row_number,
                column=column,
                value=value,
            )

    for column in worksheet.columns:
        worksheet.column_dimensions[
            column[0].column_letter
        ].width = 22

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return Response(
        content=output.getvalue(),
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition":
                'attachment; filename="FoodCare_Storage_Compliance_Report.xlsx"'
        },
    )


@router.get("/compliance-report/pdf")
def download_storage_compliance_pdf(
    limit: int = Query(500, ge=1, le=1000),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    data = get_storage_compliance_report(limit, current_user, db)

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import (
        SimpleDocTemplate,
        Table,
        TableStyle,
        Paragraph,
    )
    from reportlab.lib.styles import getSampleStyleSheet

    output = BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
    )

    styles = getSampleStyleSheet()

    elements = [
        Paragraph(
            "FoodCare - Storage Compliance Report",
            styles["Title"],
        ),
        Paragraph(
            f"Total Readings: {data['total_readings']} | "
            f"Compliant: {data['compliant_readings']} | "
            f"Non-Compliant: {data['non_compliant_readings']} | "
            f"Compliance: {data['compliance_percentage']}%",
            styles["Normal"],
        ),
    ]

    table_data = [[
        "Storage",
        "Temperature",
        "Humidity",
        "Air",
        "Light",
        "Compliant",
        "Recorded At",
    ]]

    for item in data["items"]:
        table_data.append([
            item["storage_location"],
            f"{item['temperature_c']} °C",
            f"{item['humidity_pct']}%",
            item["air_circulation"] or "",
            item["light_exposure"] or "",
            str(item["is_compliant"]),
            str(item["recorded_at"]),
        ])

    table = Table(table_data, repeatRows=1)

    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))

    elements.append(table)
    document.build(elements)

    output.seek(0)

    return Response(
        content=output.getvalue(),
        media_type="application/pdf",
        headers={
            "Content-Disposition":
                'attachment; filename="FoodCare_Storage_Compliance_Report.pdf"'
        },
    )