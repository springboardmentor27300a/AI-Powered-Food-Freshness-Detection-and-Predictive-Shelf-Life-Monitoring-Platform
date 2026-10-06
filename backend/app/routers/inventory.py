"""
/api/inventory — combined read view + dashboard summary.
"""

from fastapi import APIRouter, Depends, Query
from sqlalchemy.orm import Session, joinedload

from app.dependencies.auth import get_current_user
from app.dependencies.db import get_db
from app.models.batch import Batch
from app.models.user import User
from app.schemas.batch import BatchOut, InventorySummary
from app.services.inventory_service import get_inventory_summary
from io import BytesIO
import csv
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font

router = APIRouter(prefix="/api/inventory", tags=["inventory"])


@router.get("", response_model=list[BatchOut])
def get_inventory(
    limit: int = Query(50, ge=1, le=200),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Combined food+batch view: batches with their food item preloaded."""
    batches = (
        db.query(Batch)
        .options(joinedload(Batch.food_item))
        .order_by(Batch.expiry_date.asc())
        .limit(limit)
        .all()
    )

    return batches


@router.get("/summary", response_model=InventorySummary)
def get_summary(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    return get_inventory_summary(db)


@router.get("/quality-report")
def get_inventory_quality_report(
    limit: int = Query(200, ge=1, le=500),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """
    Inventory quality report based on the current batch status,
    quantity and expiry information.
    """

    batches = (
        db.query(Batch)
        .options(joinedload(Batch.food_item))
        .order_by(Batch.expiry_date.asc())
        .limit(limit)
        .all()
    )

    rows = []

    for batch in batches:
        batch.status = batch.compute_status()

        rows.append(
            {
                "batch_id": str(batch.id),
                "batch_code": batch.batch_code,
                "food_item_id": str(batch.food_item_id),
                "food_name": batch.food_item.name,
                "category": (
                    batch.food_item.category.value
                    if hasattr(batch.food_item.category, "value")
                    else str(batch.food_item.category)
                ),
                "quantity": batch.quantity,
                "unit": batch.unit.value,
                "expiry_date": batch.expiry_date,
                "status": batch.status.value,
                "is_available": batch.is_available,
                "storage_location": batch.storage_location,
            }
        )

    return {
        "total_batches": len(rows),
        "available_batches": sum(
            1 for row in rows if row["status"] == "available"
        ),
        "low_stock_batches": sum(
            1 for row in rows if row["status"] == "low_stock"
        ),
        "near_expiry_batches": sum(
            1 for row in rows if row["status"] == "near_expiry"
        ),
        "expired_batches": sum(
            1 for row in rows if row["status"] == "expired"
        ),
        "out_of_stock_batches": sum(
            1 for row in rows if row["status"] == "out_of_stock"
        ),
        "items": rows,
    }
@router.get("/quality-report/csv")
def download_inventory_quality_csv(
    limit: int = Query(200, ge=1, le=500),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    data = get_inventory_quality_report(limit, current_user, db)

    output = BytesIO()
    text_output = output

    import io
    csv_buffer = io.StringIO()

    writer = csv.writer(csv_buffer)

    writer.writerow([
        "Batch Code",
        "Food Name",
        "Category",
        "Quantity",
        "Unit",
        "Expiry Date",
        "Status",
        "Available",
        "Storage Location",
    ])

    for item in data["items"]:
        writer.writerow([
            item["batch_code"],
            item["food_name"],
            item["category"],
            item["quantity"],
            item["unit"],
            item["expiry_date"],
            item["status"],
            item["is_available"],
            item["storage_location"],
        ])

    return Response(
        content=csv_buffer.getvalue(),
        media_type="text/csv",
        headers={
            "Content-Disposition":
                'attachment; filename="FoodCare_Inventory_Quality_Report.csv"'
        },
    )


@router.get("/quality-report/excel")
def download_inventory_quality_excel(
    limit: int = Query(200, ge=1, le=500),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    data = get_inventory_quality_report(limit, current_user, db)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Inventory Quality"

    headers = [
        "Batch Code",
        "Food Name",
        "Category",
        "Quantity",
        "Unit",
        "Expiry Date",
        "Status",
        "Available",
        "Storage Location",
    ]

    for column, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=column, value=header)
        cell.font = Font(bold=True)

    for row_number, item in enumerate(data["items"], start=2):
        values = [
            item["batch_code"],
            item["food_name"],
            item["category"],
            item["quantity"],
            item["unit"],
            str(item["expiry_date"]),
            item["status"],
            item["is_available"],
            item["storage_location"],
        ]

        for column, value in enumerate(values, start=1):
            worksheet.cell(
                row=row_number,
                column=column,
                value=value,
            )

    for column in worksheet.columns:
        worksheet.column_dimensions[column[0].column_letter].width = 20

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return Response(
        content=output.getvalue(),
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition":
                'attachment; filename="FoodCare_Inventory_Quality_Report.xlsx"'
        },
    )
@router.get("/quality-report/pdf")
def download_inventory_quality_pdf(
    limit: int = Query(200, ge=1, le=500),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    data = get_inventory_quality_report(limit, current_user, db)

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import (
        SimpleDocTemplate,
        Table,
        TableStyle,
        Paragraph,
    )
    from reportlab.lib.styles import getSampleStyleSheet

    output = BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
    )

    styles = getSampleStyleSheet()

    elements = [
        Paragraph(
            "FoodCare - Inventory Quality Report",
            styles["Title"],
        ),
        Paragraph(
            f"Total Batches: {data['total_batches']} | "
            f"Available: {data['available_batches']} | "
            f"Near Expiry: {data['near_expiry_batches']} | "
            f"Expired: {data['expired_batches']}",
            styles["Normal"],
        ),
    ]

    table_data = [[
        "Batch",
        "Food",
        "Category",
        "Quantity",
        "Expiry",
        "Status",
        "Available",
        "Storage",
    ]]

    for item in data["items"]:
        table_data.append([
            item["batch_code"],
            item["food_name"],
            item["category"],
            f"{item['quantity']} {item['unit']}",
            str(item["expiry_date"]),
            item["status"],
            str(item["is_available"]),
            item["storage_location"] or "",
        ])

    table = Table(table_data, repeatRows=1)

    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))

    elements.append(table)
    document.build(elements)

    output.seek(0)

    return Response(
        content=output.getvalue(),
        media_type="application/pdf",
        headers={
            "Content-Disposition":
                'attachment; filename="FoodCare_Inventory_Quality_Report.pdf"'
        },
    )