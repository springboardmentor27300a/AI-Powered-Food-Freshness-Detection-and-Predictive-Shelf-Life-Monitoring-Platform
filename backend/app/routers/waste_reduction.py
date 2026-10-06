"""
/api/waste-reduction — waste reduction reporting.
"""

from datetime import date
from io import BytesIO
import csv
import io

from fastapi import APIRouter, Depends, Query
from fastapi.responses import Response
from sqlalchemy.orm import Session, joinedload

from openpyxl import Workbook
from openpyxl.styles import Font

from app.dependencies.auth import get_current_user
from app.dependencies.db import get_db
from app.models.batch import Batch
from app.models.recommendation import (
    Recommendation,
    RecommendationType,
)
from app.models.user import User

router = APIRouter(
    prefix="/api/waste-reduction",
    tags=["waste-reduction"],
)


@router.get("/report")
def get_waste_reduction_report(
    limit: int = Query(200, ge=1, le=500),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """
    Report batches that are at risk of waste based on expiry
    and existing waste-reduction recommendations.
    """

    batches = (
        db.query(Batch)
        .options(joinedload(Batch.food_item))
        .order_by(Batch.expiry_date.asc())
        .limit(limit)
        .all()
    )

    items = []

    for batch in batches:
        days_to_expiry = (batch.expiry_date - date.today()).days

        if days_to_expiry <= 5 and batch.quantity > 0:
            recommendation = (
                db.query(Recommendation)
                .filter(
                    Recommendation.batch_id == batch.id,
                    Recommendation.type
                    == RecommendationType.WASTE_REDUCTION,
                )
                .order_by(Recommendation.created_at.desc())
                .first()
            )

            items.append(
                {
                    "batch_id": str(batch.id),
                    "batch_code": batch.batch_code,
                    "food_name": batch.food_item.name,
                    "quantity": batch.quantity,
                    "unit": batch.unit.value,
                    "expiry_date": batch.expiry_date,
                    "days_to_expiry": days_to_expiry,
                    "storage_location": batch.storage_location,
                    "recommendation": (
                        {
                            "priority": recommendation.priority.value,
                            "title": recommendation.title,
                            "message": recommendation.message,
                        }
                        if recommendation
                        else None
                    ),
                }
            )

    return {
        "total_at_risk_batches": len(items),
        "total_at_risk_quantity": sum(
            item["quantity"] for item in items
        ),
        "items": items,
    }


@router.get("/report/csv")
def download_waste_reduction_csv(
    limit: int = Query(200, ge=1, le=500),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    data = get_waste_reduction_report(limit, current_user, db)

    csv_buffer = io.StringIO()
    writer = csv.writer(csv_buffer)

    writer.writerow([
        "Batch Code",
        "Food Name",
        "Quantity",
        "Unit",
        "Expiry Date",
        "Days to Expiry",
        "Storage Location",
        "Recommendation",
    ])

    for item in data["items"]:
        recommendation = item["recommendation"]

        writer.writerow([
            item["batch_code"],
            item["food_name"],
            item["quantity"],
            item["unit"],
            item["expiry_date"],
            item["days_to_expiry"],
            item["storage_location"],
            recommendation["message"] if recommendation else "",
        ])

    return Response(
        content=csv_buffer.getvalue(),
        media_type="text/csv",
        headers={
            "Content-Disposition":
                'attachment; filename="FoodCare_Waste_Reduction_Report.csv"'
        },
    )


@router.get("/report/excel")
def download_waste_reduction_excel(
    limit: int = Query(200, ge=1, le=500),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    data = get_waste_reduction_report(limit, current_user, db)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Waste Reduction"

    headers = [
        "Batch Code",
        "Food Name",
        "Quantity",
        "Unit",
        "Expiry Date",
        "Days to Expiry",
        "Storage Location",
        "Priority",
        "Recommendation",
    ]

    for column, header in enumerate(headers, start=1):
        cell = worksheet.cell(
            row=1,
            column=column,
            value=header,
        )
        cell.font = Font(bold=True)

    for row_number, item in enumerate(
        data["items"],
        start=2,
    ):
        recommendation = item["recommendation"]

        values = [
            item["batch_code"],
            item["food_name"],
            item["quantity"],
            item["unit"],
            str(item["expiry_date"]),
            item["days_to_expiry"],
            item["storage_location"],
            recommendation["priority"]
            if recommendation
            else "",
            recommendation["message"]
            if recommendation
            else "",
        ]

        for column, value in enumerate(
            values,
            start=1,
        ):
            worksheet.cell(
                row=row_number,
                column=column,
                value=value,
            )

    for column in worksheet.columns:
        worksheet.column_dimensions[
            column[0].column_letter
        ].width = 22

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return Response(
        content=output.getvalue(),
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition":
                'attachment; filename="FoodCare_Waste_Reduction_Report.xlsx"'
        },
    )


@router.get("/report/pdf")
def download_waste_reduction_pdf(
    limit: int = Query(200, ge=1, le=500),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    data = get_waste_reduction_report(limit, current_user, db)

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import (
        SimpleDocTemplate,
        Table,
        TableStyle,
        Paragraph,
    )
    from reportlab.lib.styles import getSampleStyleSheet

    output = BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
    )

    styles = getSampleStyleSheet()

    elements = [
        Paragraph(
            "FoodCare - Waste Reduction Report",
            styles["Title"],
        ),
        Paragraph(
            f"Total At-Risk Batches: "
            f"{data['total_at_risk_batches']}",
            styles["Normal"],
        ),
    ]

    table_data = [[
        "Batch",
        "Food",
        "Quantity",
        "Expiry",
        "Days Left",
        "Storage",
        "Recommendation",
    ]]

    for item in data["items"]:
        recommendation = item["recommendation"]

        table_data.append([
            item["batch_code"],
            item["food_name"],
            f"{item['quantity']} {item['unit']}",
            str(item["expiry_date"]),
            str(item["days_to_expiry"]),
            item["storage_location"] or "",
            recommendation["message"]
            if recommendation
            else "",
        ])

    table = Table(
        table_data,
        repeatRows=1,
    )

    table.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.lightgrey,
            ),
            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold",
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.5,
                colors.grey,
            ),
            (
                "FONTSIZE",
                (0, 0),
                (-1, -1),
                8,
            ),
            (
                "VALIGN",
                (0, 0),
                (-1, -1),
                "MIDDLE",
            ),
        ])
    )

    elements.append(table)
    document.build(elements)

    output.seek(0)

    return Response(
        content=output.getvalue(),
        media_type="application/pdf",
        headers={
            "Content-Disposition":
                'attachment; filename="FoodCare_Waste_Reduction_Report.pdf"'
        },
    )