"""
reports.py — FreshGuard Professional PDF Report Generator
Generates PDF reports for food freshness assessment using ReportLab.
"""

import json
from io import BytesIO
from datetime import datetime, timezone, date
from typing import Any, Dict, List, Optional
from xml.sax.saxutils import escape

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    HRFlowable,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

# FreshGuard Color System
PRIMARY = colors.HexColor("#0f172a")      # Slate 900
EMERALD = colors.HexColor("#059669")      # Emerald 600
TEXT_MAIN = colors.HexColor("#1e293b")    # Slate 800
TEXT_MUTED = colors.HexColor("#64748b")   # Slate 500
BG_LIGHT = colors.HexColor("#f8fafc")     # Slate 50
BORDER_COLOR = colors.HexColor("#e2e8f0") # Slate 200

# Classification status badge colors
STATUS_COLORS: Dict[str, Dict[str, Any]] = {
    "Fresh": {"text": colors.HexColor("#166534"), "bg": colors.HexColor("#dcfce7"), "border": colors.HexColor("#86efac")},
    "Good": {"text": colors.HexColor("#3f6212"), "bg": colors.HexColor("#ecfccb"), "border": colors.HexColor("#bef264")},
    "Acceptable": {"text": colors.HexColor("#854d0e"), "bg": colors.HexColor("#fef9c3"), "border": colors.HexColor("#fde047")},
    "Near Spoilage": {"text": colors.HexColor("#9a3412"), "bg": colors.HexColor("#ffedd5"), "border": colors.HexColor("#fdba74")},
    "Spoiled": {"text": colors.HexColor("#991b1b"), "bg": colors.HexColor("#fee2e2"), "border": colors.HexColor("#fca5a5")},
}


def _safe_json_loads(data: Any) -> Dict[str, Any]:
    if not data:
        return {}
    if isinstance(data, dict):
        return data
    try:
        return json.loads(str(data))
    except Exception:
        return {}


def generate_freshness_pdf_report(analysis: Any, user: Any) -> bytes:
    """
    Generates a PDF freshness report for the given FreshnessAnalysis and User.
    Returns raw PDF bytes.
    """
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36,
    )

    styles = getSampleStyleSheet()

    # Custom typography styles
    style_brand = ParagraphStyle(
        "BrandName",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        textColor=EMERALD,
    )
    style_tagline = ParagraphStyle(
        "BrandTagline",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=11,
        textColor=TEXT_MUTED,
    )
    style_title = ParagraphStyle(
        "DocTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=20,
        textColor=PRIMARY,
        alignment=2, # Right aligned header title
    )
    style_subtitle = ParagraphStyle(
        "DocSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=12,
        textColor=TEXT_MUTED,
        alignment=2,
    )
    style_section_h2 = ParagraphStyle(
        "SectionH2",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=16,
        textColor=PRIMARY,
        spaceBefore=8,
        spaceAfter=4,
    )
    style_body = ParagraphStyle(
        "BodyTextCustom",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=13,
        textColor=TEXT_MAIN,
    )
    style_body_bold = ParagraphStyle(
        "BodyTextBold",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=9,
        leading=13,
        textColor=TEXT_MAIN,
    )
    style_disclaimer = ParagraphStyle(
        "DisclaimerText",
        parent=styles["Normal"],
        fontName="Helvetica-Oblique",
        fontSize=8,
        leading=11,
        textColor=TEXT_MUTED,
    )

    story = []

    # -------------------------------------------------------------------------
    # A. Header Block
    # -------------------------------------------------------------------------
    report_id = f"FG-RPT-{analysis.id:06d}" if getattr(analysis, "id", None) else "FG-RPT-000000"
    now_str = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")

    header_left = [
        Paragraph("🌿 FreshGuard", style_brand),
        Paragraph("AI-Powered Food Freshness Platform", style_tagline),
    ]
    header_right = [
        Paragraph("Food Freshness Assessment Report", style_title),
        Paragraph(f"<b>Report ID:</b> {report_id} &bull; <b>Generated:</b> {now_str}", style_subtitle),
    ]

    header_table = Table([[header_left, header_right]], colWidths=[3.25 * inch, 4.25 * inch])
    header_table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
    ]))
    story.append(header_table)
    story.append(Spacer(1, 10))
    story.append(HRFlowable(width="100%", thickness=1.5, color=EMERALD, spaceBefore=0, spaceAfter=12))

    # -------------------------------------------------------------------------
    # B. Metadata / Food Information
    # -------------------------------------------------------------------------
    user_name = getattr(user, "full_name", "Authenticated User")
    user_email = getattr(user, "email", "N/A")
    user_role = getattr(user, "role", "consumer").replace("_", " ").title()
    item_name = getattr(analysis, "item_name", None) or "Unspecified Food Item"
    created_at = getattr(analysis, "created_at", None)
    scan_date_str = created_at.strftime("%Y-%m-%d %H:%M UTC") if created_at else now_str

    info_data = [
        [
            Paragraph("<b>Food Type / Classification:</b>", style_body),
            Paragraph(str(item_name), style_body_bold),
            Paragraph("<b>Analysis ID:</b>", style_body),
            Paragraph(str(getattr(analysis, "id", "N/A")), style_body),
        ],
        [
            Paragraph("<b>User / Inspector:</b>", style_body),
            Paragraph(f"{user_name} ({user_role})", style_body),
            Paragraph("<b>User Email:</b>", style_body),
            Paragraph(user_email, style_body),
        ],
        [
            Paragraph("<b>Raw Model Class:</b>", style_body),
            Paragraph(str(getattr(analysis, "predicted_class", "N/A")), style_body),
            Paragraph("<b>Scan Timestamp:</b>", style_body),
            Paragraph(scan_date_str, style_body),
        ],
    ]

    info_table = Table(info_data, colWidths=[1.8 * inch, 2.2 * inch, 1.3 * inch, 2.2 * inch])
    info_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), BG_LIGHT),
        ("BOX", (0, 0), (-1, -1), 0.5, BORDER_COLOR),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, BORDER_COLOR),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
    ]))
    story.append(info_table)
    story.append(Spacer(1, 14))

    # -------------------------------------------------------------------------
    # C. AI Assessment Results Summary
    # -------------------------------------------------------------------------
    story.append(Paragraph("AI Assessment Overview", style_section_h2))

    classification = getattr(analysis, "classification", "Acceptable")
    badge_cfg = STATUS_COLORS.get(classification, STATUS_COLORS["Acceptable"])
    quality_score = int(getattr(analysis, "quality_score", 0))
    confidence_pct = f"{float(getattr(analysis, 'confidence', 0.0)) * 100:.1f}%"
    spoilage_pct = f"{float(getattr(analysis, 'spoilage_probability', 0.0)) * 100:.1f}%"

    style_status_badge = ParagraphStyle(
        "StatusBadge",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=14,
        leading=16,
        textColor=badge_cfg["text"],
        alignment=1, # Center
    )
    style_metric_val = ParagraphStyle(
        "MetricVal",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=15,
        leading=18,
        textColor=PRIMARY,
        alignment=1,
    )
    style_metric_lbl = ParagraphStyle(
        "MetricLbl",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8,
        leading=10,
        textColor=TEXT_MUTED,
        alignment=1,
    )

    kpi_data = [
        [
            Paragraph("FRESHNESS STATUS", style_metric_lbl),
            Paragraph("QUALITY SCORE", style_metric_lbl),
            Paragraph("MODEL CONFIDENCE", style_metric_lbl),
            Paragraph("SPOILAGE PROBABILITY", style_metric_lbl),
        ],
        [
            Paragraph(classification.upper(), style_status_badge),
            Paragraph(f"{quality_score} / 100", style_metric_val),
            Paragraph(confidence_pct, style_metric_val),
            Paragraph(spoilage_pct, style_metric_val),
        ],
    ]

    kpi_table = Table(kpi_data, colWidths=[1.9 * inch, 1.8 * inch, 1.9 * inch, 1.9 * inch])
    kpi_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, 1), badge_cfg["bg"]),
        ("BOX", (0, 0), (0, 1), 1, badge_cfg["border"]),
        ("BACKGROUND", (1, 0), (-1, -1), BG_LIGHT),
        ("BOX", (1, 0), (-1, -1), 0.5, BORDER_COLOR),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, BORDER_COLOR),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(kpi_table)
    story.append(Spacer(1, 14))

    # -------------------------------------------------------------------------
    # D. Visual Analysis Breakdown (OpenCV Metrics)
    # -------------------------------------------------------------------------
    story.append(Paragraph("Visual Analysis Metrics (OpenCV Computer Vision)", style_section_h2))

    color_metrics = _safe_json_loads(getattr(analysis, "color_metrics", None))
    texture_metrics = _safe_json_loads(getattr(analysis, "texture_metrics", None))

    dom_hex = color_metrics.get("dominant_color_hex", "N/A")
    discoloration = f"{float(color_metrics.get('discoloration_index', 0.0)) * 100:.1f}%"
    browning = f"{float(color_metrics.get('browning_index', 0.0)) * 100:.1f}%"
    dark_spots = f"{float(color_metrics.get('dark_spot_ratio', 0.0)) * 100:.1f}%"
    saturation = f"{float(color_metrics.get('mean_saturation', 0.0)) * 100:.1f}%"
    brightness = f"{float(color_metrics.get('mean_brightness', 0.0)):.1f} / 255"

    uniformity = f"{float(texture_metrics.get('texture_uniformity', 0.0)) * 100:.1f}%"
    roughness = f"{float(texture_metrics.get('surface_roughness', 0.0)):.2f}"
    contrast = f"{float(texture_metrics.get('local_contrast', 0.0)):.2f}"

    metrics_headers = [
        Paragraph("<b>Colour & Surface Feature</b>", style_body_bold),
        Paragraph("<b>Measured Value</b>", style_body_bold),
        Paragraph("<b>Texture & Structural Metric</b>", style_body_bold),
        Paragraph("<b>Measured Value</b>", style_body_bold),
    ]

    metrics_rows = [
        metrics_headers,
        [
            Paragraph("Dominant Colour Hex", style_body),
            Paragraph(f"<code>{dom_hex}</code>", style_body),
            Paragraph("Texture Uniformity", style_body),
            Paragraph(uniformity, style_body),
        ],
        [
            Paragraph("Discoloration Index", style_body),
            Paragraph(discoloration, style_body),
            Paragraph("Surface Roughness", style_body),
            Paragraph(roughness, style_body),
        ],
        [
            Paragraph("Browning Index", style_body),
            Paragraph(browning, style_body),
            Paragraph("Local Contrast Variance", style_body),
            Paragraph(contrast, style_body),
        ],
        [
            Paragraph("Dark Spot Ratio", style_body),
            Paragraph(dark_spots, style_body),
            Paragraph("Feature Status", style_body),
            Paragraph("Extracted cleanly via OpenCV", style_body),
        ],
        [
            Paragraph("Mean Saturation", style_body),
            Paragraph(saturation, style_body),
            Paragraph("Mean Brightness", style_body),
            Paragraph(brightness, style_body),
        ],
    ]

    metrics_table = Table(metrics_rows, colWidths=[2.1 * inch, 1.65 * inch, 2.1 * inch, 1.65 * inch])
    metrics_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), BG_LIGHT),
        ("LINEBELOW", (0, 0), (-1, 0), 1, PRIMARY),
        ("BOX", (0, 0), (-1, -1), 0.5, BORDER_COLOR),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, BORDER_COLOR),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(metrics_table)
    story.append(Spacer(1, 14))

    # -------------------------------------------------------------------------
    # E. Interpretation Section
    # -------------------------------------------------------------------------
    story.append(Paragraph("Assessment Interpretation & Methodology", style_section_h2))

    interp_text = (
        f"The visual analysis for <b>{item_name}</b> resulted in a quality score of <b>{quality_score}/100</b> "
        f"with a predicted freshness category of <b>{classification}</b> (Confidence: {confidence_pct}, "
        f"Spoilage Risk: {spoilage_pct}).<br/>"
        "<b>Visual Assessment Note:</b> This assessment uses non-invasive visual feature extraction "
        "including HSV colour space distribution, browning index metrics, surface texture uniformity, "
        "and MobileNetV2 deep neural network classification. The visual score is a transparent weighted synthesis "
        "of predicted class probability, surface discoloration, and texture smoothness."
    )
    story.append(Paragraph(interp_text, style_body))
    story.append(Spacer(1, 6))

    system_note = (
        "<b>System Capability Scope:</b> FreshGuard evaluates overall visual condition, color fidelity, "
        "and surface degradation. The system does not claim to detect internal mold, internal bruising, "
        "or mechanical damage using dedicated trained detectors."
    )
    story.append(Paragraph(system_note, style_body))
    story.append(Spacer(1, 14))

    # -------------------------------------------------------------------------
    # F. Disclaimer Section
    # -------------------------------------------------------------------------
    story.append(HRFlowable(width="100%", thickness=0.5, color=BORDER_COLOR, spaceBefore=0, spaceAfter=8))
    disclaimer_text = (
        "<b>Official Disclaimer:</b> This report is an automated AI-assisted visual assessment produced by "
        "FreshGuard. It is intended for monitoring, inventory quality sorting, and shelf-life tracking purposes. "
        "It should not replace official laboratory analysis or professional food safety inspection by regulatory authorities."
    )
    story.append(Paragraph(disclaimer_text, style_disclaimer))

    # Build PDF document
    doc.build(story)
    pdf_bytes = buffer.getvalue()
    buffer.close()
    return pdf_bytes


# ============================================================
# M4 — Additional Report Generators
# ============================================================

# Synthetic model categories
SYNTHETIC_CATEGORIES = {"vegetables", "meat_poultry", "bakery", "packaged_foods", "beverages"}
REAL_CATEGORIES = {"tomato", "paneer", "seafood"}


def _get_data_type_label(food_type: str) -> str:
    """Return the data type label for a food type."""
    if food_type.lower() in SYNTHETIC_CATEGORIES:
        return "Synthetic Dataset (Demonstration)"
    elif food_type.lower() in REAL_CATEGORIES:
        return "Real/Traceable"
    return "Unknown"


def _get_risk_color(risk_level: str) -> colors.HexColor:
    """Return color for risk level."""
    risk_colors = {
        "low": colors.HexColor("#166534"),
        "moderate": colors.HexColor("#854d0e"),
        "high": colors.HexColor("#9a3412"),
        "critical": colors.HexColor("#991b1b"),
    }
    return risk_colors.get(risk_level.lower(), colors.HexColor("#64748b"))


def generate_shelf_life_pdf_report(predictions: List[Any], user: Any) -> bytes:
    """Generate a PDF shelf-life report from M3 prediction data."""
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()

    style_brand = ParagraphStyle("BrandName", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=20, leading=24, textColor=EMERALD)
    style_tagline = ParagraphStyle("BrandTagline", parent=styles["Normal"], fontName="Helvetica", fontSize=9, leading=11, textColor=TEXT_MUTED)
    style_title = ParagraphStyle("DocTitle", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=16, leading=20, textColor=PRIMARY, alignment=2)
    style_subtitle = ParagraphStyle("DocSubtitle", parent=styles["Normal"], fontName="Helvetica", fontSize=9, leading=12, textColor=TEXT_MUTED, alignment=2)
    style_section = ParagraphStyle("SectionH2", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=12, leading=16, textColor=PRIMARY, spaceBefore=8, spaceAfter=4)
    style_body = ParagraphStyle("BodyText", parent=styles["Normal"], fontName="Helvetica", fontSize=9, leading=13, textColor=TEXT_MAIN)
    style_body_bold = ParagraphStyle("BodyBold", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=9, leading=13, textColor=TEXT_MAIN)
    style_disclaimer = ParagraphStyle("Disclaimer", parent=styles["Normal"], fontName="Helvetica-Oblique", fontSize=8, leading=11, textColor=TEXT_MUTED)

    story = []
    now_str = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")

    # Header
    header_left = [Paragraph("FreshGuard", style_brand), Paragraph("AI-Powered Food Freshness Platform", style_tagline)]
    header_right = [Paragraph("Shelf-Life Prediction Report", style_title), Paragraph(f"<b>Generated:</b> {now_str}", style_subtitle)]
    header_table = Table([[header_left, header_right]], colWidths=[3.25 * inch, 4.25 * inch])
    header_table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
    story.append(header_table)
    story.append(Spacer(1, 10))
    story.append(HRFlowable(width="100%", thickness=1.5, color=EMERALD, spaceBefore=0, spaceAfter=12))

    # User info
    user_name = getattr(user, "full_name", "Authenticated User")
    user_role = getattr(user, "role", "consumer").replace("_", " ").title()
    story.append(Paragraph(f"<b>User:</b> {user_name} ({user_role})", style_body))
    story.append(Spacer(1, 8))

    # Summary metrics
    total = len(predictions)
    expired = sum(1 for p in predictions if p.remaining_shelf_life_days <= 0)
    expiring_soon = sum(1 for p in predictions if 0 < p.remaining_shelf_life_days <= 2)
    high_risk = sum(1 for p in predictions if p.risk_level in ("high", "critical"))

    story.append(Paragraph("Summary", style_section))
    summary_data = [
        [Paragraph("<b>Total Predictions</b>", style_body), Paragraph(str(total), style_body_bold)],
        [Paragraph("<b>Expired</b>", style_body), Paragraph(str(expired), style_body_bold)],
        [Paragraph("<b>Expiring Soon (1-2 days)</b>", style_body), Paragraph(str(expiring_soon), style_body_bold)],
        [Paragraph("<b>High Risk</b>", style_body), Paragraph(str(high_risk), style_body_bold)],
    ]
    summary_table = Table(summary_data, colWidths=[3 * inch, 2 * inch])
    summary_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), BG_LIGHT), ("BOX", (0, 0), (-1, -1), 0.5, BORDER_COLOR), ("INNERGRID", (0, 0), (-1, -1), 0.5, BORDER_COLOR), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5), ("LEFTPADDING", (0, 0), (-1, -1), 8)]))
    story.append(summary_table)
    story.append(Spacer(1, 14))

    # Prediction details table
    story.append(Paragraph("Prediction Details", style_section))
    if predictions:
        table_data = [[Paragraph("<b>Food Type</b>", style_body_bold), Paragraph("<b>Data Type</b>", style_body_bold), Paragraph("<b>Remaining Days</b>", style_body_bold), Paragraph("<b>Expiry Date</b>", style_body_bold), Paragraph("<b>Risk</b>", style_body_bold), Paragraph("<b>Model</b>", style_body_bold)]]
        for pred in predictions[:20]:
            table_data.append([
                Paragraph(str(pred.food_type).title(), style_body),
                Paragraph(_get_data_type_label(pred.food_type), style_body),
                Paragraph(str(pred.remaining_shelf_life_days), style_body),
                Paragraph(str(pred.predicted_expiry_date), style_body),
                Paragraph(pred.risk_level.upper(), style_body),
                Paragraph(str(pred.model_type), style_body),
            ])
        pred_table = Table(table_data, colWidths=[1.1 * inch, 1.3 * inch, 1 * inch, 1.1 * inch, 0.8 * inch, 1.2 * inch])
        pred_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), BG_LIGHT), ("LINEBELOW", (0, 0), (-1, 0), 1, PRIMARY), ("BOX", (0, 0), (-1, -1), 0.5, BORDER_COLOR), ("INNERGRID", (0, 0), (-1, -1), 0.5, BORDER_COLOR), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
        story.append(pred_table)
    else:
        story.append(Paragraph("No shelf-life predictions available.", style_body))
    story.append(Spacer(1, 14))

    # Disclaimer
    story.append(HRFlowable(width="100%", thickness=0.5, color=BORDER_COLOR, spaceBefore=0, spaceAfter=8))
    story.append(Paragraph("<b>Disclaimer:</b> Shelf-life predictions are generated by ML models. Real/Traceable models (Tomato, Paneer, Seafood) are trained on published research datasets. Synthetic Dataset (Demonstration) models are for demonstration purposes only and are not real-world validated.", style_disclaimer))

    doc.build(story)
    pdf_bytes = buffer.getvalue()
    buffer.close()
    return pdf_bytes


def generate_inventory_quality_pdf_report(inventory: List[Any], predictions: List[Any], user: Any) -> bytes:
    """Generate a PDF inventory quality report."""
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()

    style_brand = ParagraphStyle("BrandName", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=20, leading=24, textColor=EMERALD)
    style_tagline = ParagraphStyle("BrandTagline", parent=styles["Normal"], fontName="Helvetica", fontSize=9, leading=11, textColor=TEXT_MUTED)
    style_title = ParagraphStyle("DocTitle", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=16, leading=20, textColor=PRIMARY, alignment=2)
    style_subtitle = ParagraphStyle("DocSubtitle", parent=styles["Normal"], fontName="Helvetica", fontSize=9, leading=12, textColor=TEXT_MUTED, alignment=2)
    style_section = ParagraphStyle("SectionH2", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=12, leading=16, textColor=PRIMARY, spaceBefore=8, spaceAfter=4)
    style_body = ParagraphStyle("BodyText", parent=styles["Normal"], fontName="Helvetica", fontSize=9, leading=13, textColor=TEXT_MAIN)
    style_body_bold = ParagraphStyle("BodyBold", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=9, leading=13, textColor=TEXT_MAIN)
    style_disclaimer = ParagraphStyle("Disclaimer", parent=styles["Normal"], fontName="Helvetica-Oblique", fontSize=8, leading=11, textColor=TEXT_MUTED)

    story = []
    now_str = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
    today = date.today()

    # Header
    header_left = [Paragraph("FreshGuard", style_brand), Paragraph("AI-Powered Food Freshness Platform", style_tagline)]
    header_right = [Paragraph("Inventory Quality Report", style_title), Paragraph(f"<b>Generated:</b> {now_str}", style_subtitle)]
    header_table = Table([[header_left, header_right]], colWidths=[3.25 * inch, 4.25 * inch])
    header_table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
    story.append(header_table)
    story.append(Spacer(1, 10))
    story.append(HRFlowable(width="100%", thickness=1.5, color=EMERALD, spaceBefore=0, spaceAfter=12))

    # Summary
    total = len(inventory)
    expired = sum(1 for i in inventory if i.expiry_date and i.expiry_date < today)
    from datetime import timedelta
    soon_cutoff = today + timedelta(days=3)
    expiring_soon = sum(1 for i in inventory if i.expiry_date and today <= i.expiry_date <= soon_cutoff)
    high_risk = sum(1 for p in predictions if p.risk_level in ("high", "critical"))

    story.append(Paragraph("Inventory Summary", style_section))
    summary_data = [
        [Paragraph("<b>Total Items</b>", style_body), Paragraph(str(total), style_body_bold)],
        [Paragraph("<b>Expired</b>", style_body), Paragraph(str(expired), style_body_bold)],
        [Paragraph("<b>Expiring Soon</b>", style_body), Paragraph(str(expiring_soon), style_body_bold)],
        [Paragraph("<b>High Risk Predictions</b>", style_body), Paragraph(str(high_risk), style_body_bold)],
    ]
    summary_table = Table(summary_data, colWidths=[3 * inch, 2 * inch])
    summary_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), BG_LIGHT), ("BOX", (0, 0), (-1, -1), 0.5, BORDER_COLOR), ("INNERGRID", (0, 0), (-1, -1), 0.5, BORDER_COLOR), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5), ("LEFTPADDING", (0, 0), (-1, -1), 8)]))
    story.append(summary_table)
    story.append(Spacer(1, 14))

    # Inventory table
    story.append(Paragraph("Inventory Details", style_section))
    if inventory:
        table_data = [[Paragraph("<b>Name</b>", style_body_bold), Paragraph("<b>Category</b>", style_body_bold), Paragraph("<b>Batch</b>", style_body_bold), Paragraph("<b>Qty</b>", style_body_bold), Paragraph("<b>Expiry</b>", style_body_bold), Paragraph("<b>Status</b>", style_body_bold)]]
        for item in inventory[:20]:
            days_until = (item.expiry_date - today).days if item.expiry_date else None
            if days_until is not None and days_until < 0:
                status = "Expired"
            elif days_until is not None and days_until <= 3:
                status = "Expiring Soon"
            else:
                status = "Valid"
            table_data.append([
                Paragraph(str(item.name), style_body),
                Paragraph(str(item.category), style_body),
                Paragraph(str(item.batch_number), style_body),
                Paragraph(str(item.quantity), style_body),
                Paragraph(str(item.expiry_date), style_body),
                Paragraph(status, style_body),
            ])
        inv_table = Table(table_data, colWidths=[1.3 * inch, 1.1 * inch, 1 * inch, 0.5 * inch, 1 * inch, 1 * inch])
        inv_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), BG_LIGHT), ("LINEBELOW", (0, 0), (-1, 0), 1, PRIMARY), ("BOX", (0, 0), (-1, -1), 0.5, BORDER_COLOR), ("INNERGRID", (0, 0), (-1, -1), 0.5, BORDER_COLOR), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
        story.append(inv_table)
    else:
        story.append(Paragraph("No inventory items available.", style_body))

    story.append(Spacer(1, 14))
    story.append(HRFlowable(width="100%", thickness=0.5, color=BORDER_COLOR, spaceBefore=0, spaceAfter=8))
    story.append(Paragraph("<b>Disclaimer:</b> Inventory quality metrics are derived from actual inventory data and M3 predictions.", style_disclaimer))

    doc.build(story)
    pdf_bytes = buffer.getvalue()
    buffer.close()
    return pdf_bytes


def generate_waste_reduction_pdf_report(predictions: List[Any], inventory: List[Any], user: Any) -> bytes:
    """Generate a PDF waste reduction report."""
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()

    style_brand = ParagraphStyle("BrandName", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=20, leading=24, textColor=EMERALD)
    style_tagline = ParagraphStyle("BrandTagline", parent=styles["Normal"], fontName="Helvetica", fontSize=9, leading=11, textColor=TEXT_MUTED)
    style_title = ParagraphStyle("DocTitle", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=16, leading=20, textColor=PRIMARY, alignment=2)
    style_subtitle = ParagraphStyle("DocSubtitle", parent=styles["Normal"], fontName="Helvetica", fontSize=9, leading=12, textColor=TEXT_MUTED, alignment=2)
    style_section = ParagraphStyle("SectionH2", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=12, leading=16, textColor=PRIMARY, spaceBefore=8, spaceAfter=4)
    style_body = ParagraphStyle("BodyText", parent=styles["Normal"], fontName="Helvetica", fontSize=9, leading=13, textColor=TEXT_MAIN)
    style_body_bold = ParagraphStyle("BodyBold", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=9, leading=13, textColor=TEXT_MAIN)
    style_disclaimer = ParagraphStyle("Disclaimer", parent=styles["Normal"], fontName="Helvetica-Oblique", fontSize=8, leading=11, textColor=TEXT_MUTED)

    story = []
    now_str = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
    today = date.today()

    # Header
    header_left = [Paragraph("FreshGuard", style_brand), Paragraph("AI-Powered Food Freshness Platform", style_tagline)]
    header_right = [Paragraph("Waste Reduction Report", style_title), Paragraph(f"<b>Generated:</b> {now_str}", style_subtitle)]
    header_table = Table([[header_left, header_right]], colWidths=[3.25 * inch, 4.25 * inch])
    header_table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
    story.append(header_table)
    story.append(Spacer(1, 10))
    story.append(HRFlowable(width="100%", thickness=1.5, color=EMERALD, spaceBefore=0, spaceAfter=12))

    # Potential Waste-Risk Reduction Insights
    expired_count = sum(1 for p in predictions if p.remaining_shelf_life_days <= 0)
    at_risk_count = sum(1 for p in predictions if 0 < p.remaining_shelf_life_days <= 2)
    high_risk_count = sum(1 for p in predictions if p.risk_level in ("high", "critical"))
    expired_inventory = sum(1 for i in inventory if i.expiry_date and i.expiry_date < today)

    story.append(Paragraph("Potential Waste-Risk Reduction Insights", style_section))
    story.append(Paragraph("The following metrics identify items at risk of waste. These are potential insights, not measured waste reduction.", style_body))
    story.append(Spacer(1, 8))

    summary_data = [
        [Paragraph("<b>Metric</b>", style_body_bold), Paragraph("<b>Count</b>", style_body_bold)],
        [Paragraph("Expired Predictions", style_body), Paragraph(str(expired_count), style_body_bold)],
        [Paragraph("At-Risk Predictions (1-2 days)", style_body), Paragraph(str(at_risk_count), style_body_bold)],
        [Paragraph("High-Risk Predictions", style_body), Paragraph(str(high_risk_count), style_body_bold)],
        [Paragraph("Expired Inventory Items", style_body), Paragraph(str(expired_inventory), style_body_bold)],
    ]
    summary_table = Table(summary_data, colWidths=[3.5 * inch, 1.5 * inch])
    summary_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), BG_LIGHT), ("LINEBELOW", (0, 0), (-1, 0), 1, PRIMARY), ("BOX", (0, 0), (-1, -1), 0.5, BORDER_COLOR), ("INNERGRID", (0, 0), (-1, -1), 0.5, BORDER_COLOR), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5), ("LEFTPADDING", (0, 0), (-1, -1), 8)]))
    story.append(summary_table)
    story.append(Spacer(1, 14))

    # Recommendations
    story.append(Paragraph("Priority Items for Rotation (FEFO)", style_section))
    if predictions:
        table_data = [[Paragraph("<b>Food Type</b>", style_body_bold), Paragraph("<b>Remaining Days</b>", style_body_bold), Paragraph("<b>Risk</b>", style_body_bold), Paragraph("<b>Data Type</b>", style_body_bold)]]
        for pred in sorted(predictions, key=lambda p: p.remaining_shelf_life_days)[:15]:
            table_data.append([
                Paragraph(str(pred.food_type).title(), style_body),
                Paragraph(str(pred.remaining_shelf_life_days), style_body),
                Paragraph(pred.risk_level.upper(), style_body),
                Paragraph(_get_data_type_label(pred.food_type), style_body),
            ])
        waste_table = Table(table_data, colWidths=[1.5 * inch, 1.2 * inch, 1 * inch, 1.8 * inch])
        waste_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), BG_LIGHT), ("LINEBELOW", (0, 0), (-1, 0), 1, PRIMARY), ("BOX", (0, 0), (-1, -1), 0.5, BORDER_COLOR), ("INNERGRID", (0, 0), (-1, -1), 0.5, BORDER_COLOR), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
        story.append(waste_table)
    else:
        story.append(Paragraph("No prediction data available.", style_body))

    story.append(Spacer(1, 14))
    story.append(HRFlowable(width="100%", thickness=0.5, color=BORDER_COLOR, spaceBefore=0, spaceAfter=8))
    story.append(Paragraph("<b>Disclaimer:</b> This report provides potential waste-risk reduction insights based on current prediction data. It does not claim measured waste reduction or monetary savings. Actual waste reduction requires historical tracking over time.", style_disclaimer))

    doc.build(story)
    pdf_bytes = buffer.getvalue()
    buffer.close()
    return pdf_bytes


def generate_storage_compliance_pdf_report(storage_conditions: List[Any], predictions: List[Any], user: Any) -> bytes:
    """Generate a PDF storage compliance report."""
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()

    style_brand = ParagraphStyle("BrandName", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=20, leading=24, textColor=EMERALD)
    style_tagline = ParagraphStyle("BrandTagline", parent=styles["Normal"], fontName="Helvetica", fontSize=9, leading=11, textColor=TEXT_MUTED)
    style_title = ParagraphStyle("DocTitle", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=16, leading=20, textColor=PRIMARY, alignment=2)
    style_subtitle = ParagraphStyle("DocSubtitle", parent=styles["Normal"], fontName="Helvetica", fontSize=9, leading=12, textColor=TEXT_MUTED, alignment=2)
    style_section = ParagraphStyle("SectionH2", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=12, leading=16, textColor=PRIMARY, spaceBefore=8, spaceAfter=4)
    style_body = ParagraphStyle("BodyText", parent=styles["Normal"], fontName="Helvetica", fontSize=9, leading=13, textColor=TEXT_MAIN)
    style_body_bold = ParagraphStyle("BodyBold", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=9, leading=13, textColor=TEXT_MAIN)
    style_disclaimer = ParagraphStyle("Disclaimer", parent=styles["Normal"], fontName="Helvetica-Oblique", fontSize=8, leading=11, textColor=TEXT_MUTED)

    story = []
    now_str = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")

    # Header
    header_left = [Paragraph("FreshGuard", style_brand), Paragraph("AI-Powered Food Freshness Platform", style_tagline)]
    header_right = [Paragraph("Storage Compliance Report", style_title), Paragraph(f"<b>Generated:</b> {now_str}", style_subtitle)]
    header_table = Table([[header_left, header_right]], colWidths=[3.25 * inch, 4.25 * inch])
    header_table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
    story.append(header_table)
    story.append(Spacer(1, 10))
    story.append(HRFlowable(width="100%", thickness=1.5, color=EMERALD, spaceBefore=0, spaceAfter=12))

    # Compliance summary
    compliant = sum(1 for p in predictions if p.storage_compliance == "compliant")
    non_compliant = sum(1 for p in predictions if p.storage_compliance == "needs_attention")

    story.append(Paragraph("Compliance Summary", style_section))
    summary_data = [
        [Paragraph("<b>Compliant Records</b>", style_body), Paragraph(str(compliant), style_body_bold)],
        [Paragraph("<b>Needs Attention</b>", style_body), Paragraph(str(non_compliant), style_body_bold)],
        [Paragraph("<b>Compliance Rate</b>", style_body), Paragraph(f"{round(compliant / len(predictions) * 100)}%" if predictions else "N/A", style_body_bold)],
    ]
    summary_table = Table(summary_data, colWidths=[3 * inch, 2 * inch])
    summary_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), BG_LIGHT), ("BOX", (0, 0), (-1, -1), 0.5, BORDER_COLOR), ("INNERGRID", (0, 0), (-1, -1), 0.5, BORDER_COLOR), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5), ("LEFTPADDING", (0, 0), (-1, -1), 8)]))
    story.append(summary_table)
    story.append(Spacer(1, 14))

    # Storage conditions table
    story.append(Paragraph("Storage Condition Records", style_section))
    if storage_conditions:
        table_data = [[Paragraph("<b>Temp (°C)</b>", style_body_bold), Paragraph("<b>Humidity (%)</b>", style_body_bold), Paragraph("<b>Airflow</b>", style_body_bold), Paragraph("<b>Light</b>", style_body_bold), Paragraph("<b>Packaging</b>", style_body_bold), Paragraph("<b>Duration</b>", style_body_bold)]]
        for sc in storage_conditions[:20]:
            table_data.append([
                Paragraph(str(sc.temperature_c), style_body),
                Paragraph(str(sc.humidity_pct), style_body),
                Paragraph(str(sc.air_circulation), style_body),
                Paragraph(str(sc.light_exposure), style_body),
                Paragraph(str(sc.packaging_type), style_body),
                Paragraph(f"{sc.storage_duration_days} days", style_body),
            ])
        sc_table = Table(table_data, colWidths=[0.9 * inch, 0.9 * inch, 0.8 * inch, 0.8 * inch, 1.2 * inch, 0.9 * inch])
        sc_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), BG_LIGHT), ("LINEBELOW", (0, 0), (-1, 0), 1, PRIMARY), ("BOX", (0, 0), (-1, -1), 0.5, BORDER_COLOR), ("INNERGRID", (0, 0), (-1, -1), 0.5, BORDER_COLOR), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
        story.append(sc_table)
    else:
        story.append(Paragraph("No storage condition records available.", style_body))

    story.append(Spacer(1, 14))
    story.append(HRFlowable(width="100%", thickness=0.5, color=BORDER_COLOR, spaceBefore=0, spaceAfter=8))
    story.append(Paragraph("<b>Disclaimer:</b> Storage compliance is determined by the M3 storage condition scoring algorithm. This report reflects current records only.", style_disclaimer))

    doc.build(story)
    pdf_bytes = buffer.getvalue()
    buffer.close()
    return pdf_bytes


# ============================================================
# Excel Export
# ============================================================

def generate_shelf_life_excel(predictions: List[Any]) -> bytes:
    """Generate an Excel file with shelf-life prediction data."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = "Shelf-Life Predictions"

    # Title
    ws.merge_cells("A1:H1")
    ws["A1"] = "FreshGuard — Shelf-Life Prediction Report"
    ws["A1"].font = Font(size=14, bold=True, color="059669")
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:H2")
    ws["A2"] = f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')}"
    ws["A2"].font = Font(size=9, color="64748b")
    ws["A2"].alignment = Alignment(horizontal="center")

    # Headers
    headers = ["Food Type", "Data Type", "Remaining Days", "Expiry Date", "Risk Level", "Model Type", "Confidence", "Storage Compliance"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    # Data
    for row_idx, pred in enumerate(predictions, 5):
        ws.cell(row=row_idx, column=1, value=str(pred.food_type).title())
        ws.cell(row=row_idx, column=2, value=_get_data_type_label(pred.food_type))
        ws.cell(row=row_idx, column=3, value=pred.remaining_shelf_life_days)
        ws.cell(row=row_idx, column=4, value=str(pred.predicted_expiry_date))
        ws.cell(row=row_idx, column=5, value=pred.risk_level)
        ws.cell(row=row_idx, column=6, value=pred.model_type)
        ws.cell(row=row_idx, column=7, value=pred.confidence)
        ws.cell(row=row_idx, column=8, value=pred.storage_compliance or "N/A")

    # Auto-width (skip merged cells)
    for col in ws.columns:
        max_length = 0
        first_cell = col[0]
        if not hasattr(first_cell, "column_letter"):
            continue
        column = first_cell.column_letter
        for cell in col:
            try:
                if cell.value and len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:
                pass
        ws.column_dimensions[column].width = min(max_length + 2, 30)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def generate_m2_assessment_pdf_report(analysis: Any, user: Any) -> bytes:
    """Generate a report that preserves M2 v2 abstention and safety semantics."""

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=42,
        rightMargin=42,
        topMargin=42,
        bottomMargin=42,
    )
    styles = getSampleStyleSheet()
    visual_analysis = json.loads(analysis.visual_analysis_json)
    defects = json.loads(analysis.defects_json)
    class_probabilities = (
        json.loads(analysis.class_probabilities_json)
        if analysis.class_probabilities_json else None
    )
    feature_rows = [["Feature", "Image value"]]
    for group_name in ("color", "texture"):
        for feature_name, value in visual_analysis.get(group_name, {}).items():
            feature_rows.append([
                Paragraph(escape(f"{group_name}: {feature_name}"), styles["Normal"]),
                Paragraph(escape(json.dumps(value, sort_keys=True)), styles["Normal"]),
            ])
    feature_table = Table(feature_rows, colWidths=[2.4 * inch, 4.6 * inch], repeatRows=1)
    feature_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), BG_LIGHT),
        ("TEXTCOLOR", (0, 0), (-1, 0), PRIMARY),
        ("GRID", (0, 0), (-1, -1), 0.4, BORDER_COLOR),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
    ]))
    defect_rows = [["Workflow", "Status", "Reason"]]
    for defect in defects:
        defect_rows.append([
            Paragraph(escape(str(defect["defect_type"])), styles["Normal"]),
            Paragraph(escape(str(defect["status"]).replace("_", " ")), styles["Normal"]),
            Paragraph(escape(str(defect["reason"])), styles["Normal"]),
        ])
    defect_table = Table(defect_rows, colWidths=[1.3 * inch, 1.3 * inch, 4.4 * inch], repeatRows=1)
    defect_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), BG_LIGHT),
        ("TEXTCOLOR", (0, 0), (-1, 0), PRIMARY),
        ("GRID", (0, 0), (-1, -1), 0.4, BORDER_COLOR),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
    ]))
    story = [
        Paragraph("FreshGuard M2 Image Assessment", styles["Title"]),
        Spacer(1, 10),
        Paragraph(f"Assessment ID: {analysis.id}", styles["Normal"]),
        Paragraph(f"User: {escape(str(getattr(user, 'full_name', 'Authenticated User')))}", styles["Normal"]),
        Paragraph(f"Captured: {analysis.created_at.isoformat()}", styles["Normal"]),
        Spacer(1, 10),
        Paragraph("Assessment status", styles["Heading2"]),
        Paragraph(escape(analysis.assessment_status.replace("_", " ").title()), styles["Normal"]),
        Paragraph(f"Category: {escape(analysis.category)}", styles["Normal"]),
        Paragraph(f"Food subtype: {escape(analysis.food_subtype or 'Not provided')}", styles["Normal"]),
        Paragraph(f"Reason: {escape(analysis.reason)}", styles["Normal"]),
        Paragraph(f"Model ID: {escape(analysis.model_id or 'None')}", styles["Normal"]),
        Paragraph(f"Freshness category: {escape(analysis.classification or 'Not assessed')}", styles["Normal"]),
        Paragraph(
            f"Confidence: {analysis.confidence:.4f}" if analysis.confidence is not None else "Confidence: Not assessed",
            styles["Normal"],
        ),
        Paragraph(
            f"Quality score: {analysis.quality_score}/100"
            if analysis.quality_score is not None else "Quality score: Not assessed",
            styles["Normal"],
        ),
        Paragraph(
            f"Spoilage probability: {analysis.spoilage_probability:.4f}"
            if analysis.spoilage_probability is not None else "Spoilage probability: Not assessed",
            styles["Normal"],
        ),
        Paragraph(
            "Class probabilities: "
            + (
                escape(json.dumps(class_probabilities, sort_keys=True))
                if class_probabilities is not None else "Not assessed"
            ),
            styles["Normal"],
        ),
        Paragraph("Image-derived descriptive features", styles["Heading2"]),
        Paragraph(
            f"Format: {escape(analysis.image_format)}; dimensions: "
            f"{analysis.image_width} x {analysis.image_height} pixels.",
            styles["Normal"],
        ),
        Paragraph(
            "Color and texture metrics are descriptive image features. They are not validated "
            "freshness classifications or defect detections.",
            styles["Normal"],
        ),
        feature_table,
        Paragraph("Defect workflows", styles["Heading2"]),
        defect_table,
        Spacer(1, 8),
        Paragraph(
            "A photograph cannot establish microbiological safety. Follow applicable food-safety "
            "procedures and qualified inspection.",
            styles["Italic"],
        ),
    ]
    doc.build(story)
    return buffer.getvalue()


def generate_inventory_excel(inventory: List[Any]) -> bytes:
    """Generate an Excel file with inventory data."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = "Inventory Quality"

    ws.merge_cells("A1:G1")
    ws["A1"] = "FreshGuard — Inventory Quality Report"
    ws["A1"].font = Font(size=14, bold=True, color="059669")
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:G2")
    ws["A2"] = f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')}"
    ws["A2"].font = Font(size=9, color="64748b")
    ws["A2"].alignment = Alignment(horizontal="center")

    headers = ["Name", "Category", "Batch Number", "Quantity", "Expiry Date", "Days Until Expiry", "Status"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    today = date.today()
    for row_idx, item in enumerate(inventory, 5):
        days_until = (item.expiry_date - today).days if item.expiry_date else None
        if days_until is not None and days_until < 0:
            status = "Expired"
        elif days_until is not None and days_until <= 3:
            status = "Expiring Soon"
        else:
            status = "Valid"
        ws.cell(row=row_idx, column=1, value=item.name)
        ws.cell(row=row_idx, column=2, value=item.category)
        ws.cell(row=row_idx, column=3, value=item.batch_number)
        ws.cell(row=row_idx, column=4, value=item.quantity)
        ws.cell(row=row_idx, column=5, value=str(item.expiry_date))
        ws.cell(row=row_idx, column=6, value=days_until if days_until is not None else "N/A")
        ws.cell(row=row_idx, column=7, value=status)

    for col in ws.columns:
        max_length = 0
        first_cell = col[0]
        if not hasattr(first_cell, "column_letter"):
            continue
        column = first_cell.column_letter
        for cell in col:
            try:
                if cell.value and len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:
                pass
        ws.column_dimensions[column].width = min(max_length + 2, 30)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def generate_waste_reduction_excel(predictions: List[Any]) -> bytes:
    """Generate an Excel file with waste reduction insights."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = "Waste Reduction Insights"

    ws.merge_cells("A1:E1")
    ws["A1"] = "FreshGuard — Potential Waste-Risk Reduction Insights"
    ws["A1"].font = Font(size=14, bold=True, color="059669")
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:E2")
    ws["A2"] = f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')}"
    ws["A2"].font = Font(size=9, color="64748b")
    ws["A2"].alignment = Alignment(horizontal="center")

    headers = ["Food Type", "Data Type", "Remaining Days", "Risk Level", "Priority"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    sorted_preds = sorted(predictions, key=lambda p: p.remaining_shelf_life_days)
    for row_idx, pred in enumerate(sorted_preds, 5):
        priority = "High" if pred.remaining_shelf_life_days <= 0 else "Medium" if pred.remaining_shelf_life_days <= 2 else "Low"
        ws.cell(row=row_idx, column=1, value=str(pred.food_type).title())
        ws.cell(row=row_idx, column=2, value=_get_data_type_label(pred.food_type))
        ws.cell(row=row_idx, column=3, value=pred.remaining_shelf_life_days)
        ws.cell(row=row_idx, column=4, value=pred.risk_level)
        ws.cell(row=row_idx, column=5, value=priority)

    for col in ws.columns:
        max_length = 0
        first_cell = col[0]
        if not hasattr(first_cell, "column_letter"):
            continue
        column = first_cell.column_letter
        for cell in col:
            try:
                if cell.value and len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:
                pass
        ws.column_dimensions[column].width = min(max_length + 2, 30)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def generate_storage_compliance_excel(storage_conditions: List[Any], predictions: List[Any]) -> bytes:
    """Generate an Excel file with storage compliance data."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()

    # Sheet 1: Storage Conditions
    ws1 = wb.active
    ws1.title = "Storage Conditions"
    ws1.merge_cells("A1:F1")
    ws1["A1"] = "FreshGuard — Storage Compliance Report"
    ws1["A1"].font = Font(size=14, bold=True, color="059669")
    ws1["A1"].alignment = Alignment(horizontal="center")

    headers = ["Temp (°C)", "Humidity (%)", "Airflow", "Light", "Packaging", "Duration (days)"]
    for col, header in enumerate(headers, 1):
        cell = ws1.cell(row=3, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    for row_idx, sc in enumerate(storage_conditions, 4):
        ws1.cell(row=row_idx, column=1, value=sc.temperature_c)
        ws1.cell(row=row_idx, column=2, value=sc.humidity_pct)
        ws1.cell(row=row_idx, column=3, value=sc.air_circulation)
        ws1.cell(row=row_idx, column=4, value=sc.light_exposure)
        ws1.cell(row=row_idx, column=5, value=sc.packaging_type)
        ws1.cell(row=row_idx, column=6, value=sc.storage_duration_days)

    # Sheet 2: Compliance Summary
    ws2 = wb.create_sheet("Compliance Summary")
    ws2["A1"] = "Compliance Status"
    ws2["B1"] = "Count"
    ws2["A1"].font = Font(bold=True)
    ws2["B1"].font = Font(bold=True)

    compliant = sum(1 for p in predictions if p.storage_compliance == "compliant")
    non_compliant = sum(1 for p in predictions if p.storage_compliance == "needs_attention")
    ws2["A2"] = "Compliant"
    ws2["B2"] = compliant
    ws2["A3"] = "Needs Attention"
    ws2["B3"] = non_compliant

    for col in ws1.columns:
        max_length = 0
        first_cell = col[0]
        if not hasattr(first_cell, "column_letter"):
            continue
        column = first_cell.column_letter
        for cell in col:
            try:
                if cell.value and len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:
                pass
        ws1.column_dimensions[column].width = min(max_length + 2, 30)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ============================================================
# M4 — Freshness Summary Report (Historical)
# ============================================================

def generate_freshness_summary_pdf_report(analyses: List[Any], user: Any) -> bytes:
    """Generate a PDF freshness summary report from all user's M2 analyses."""
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()

    style_brand = ParagraphStyle("BrandName", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=20, leading=24, textColor=EMERALD)
    style_tagline = ParagraphStyle("BrandTagline", parent=styles["Normal"], fontName="Helvetica", fontSize=9, leading=11, textColor=TEXT_MUTED)
    style_title = ParagraphStyle("DocTitle", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=16, leading=20, textColor=PRIMARY, alignment=2)
    style_subtitle = ParagraphStyle("DocSubtitle", parent=styles["Normal"], fontName="Helvetica", fontSize=9, leading=12, textColor=TEXT_MUTED, alignment=2)
    style_section = ParagraphStyle("SectionH2", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=12, leading=16, textColor=PRIMARY, spaceBefore=8, spaceAfter=4)
    style_body = ParagraphStyle("BodyText", parent=styles["Normal"], fontName="Helvetica", fontSize=9, leading=13, textColor=TEXT_MAIN)
    style_body_bold = ParagraphStyle("BodyBold", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=9, leading=13, textColor=TEXT_MAIN)
    style_disclaimer = ParagraphStyle("Disclaimer", parent=styles["Normal"], fontName="Helvetica-Oblique", fontSize=8, leading=11, textColor=TEXT_MUTED)

    story = []
    now_str = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")

    # Header
    header_left = [Paragraph("FreshGuard", style_brand), Paragraph("AI-Powered Food Freshness Platform", style_tagline)]
    header_right = [Paragraph("Freshness Summary Report", style_title), Paragraph(f"<b>Generated:</b> {now_str}", style_subtitle)]
    header_table = Table([[header_left, header_right]], colWidths=[3.25 * inch, 4.25 * inch])
    header_table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
    story.append(header_table)
    story.append(Spacer(1, 10))
    story.append(HRFlowable(width="100%", thickness=1.5, color=EMERALD, spaceBefore=0, spaceAfter=12))

    # User info
    user_name = getattr(user, "full_name", "Authenticated User")
    user_role = getattr(user, "role", "consumer").replace("_", " ").title()
    story.append(Paragraph(f"<b>User:</b> {user_name} ({user_role})", style_body))
    story.append(Spacer(1, 8))

    # Summary statistics
    total = len(analyses)
    if total > 0:
        avg_quality = round(sum(a.quality_score for a in analyses) / total, 1)
        avg_confidence = round(sum(a.confidence for a in analyses) / total * 100, 1)
        avg_spoilage = round(sum(a.spoilage_probability for a in analyses) / total * 100, 1)
        classification_counts = {}
        for a in analyses:
            classification_counts[a.classification] = classification_counts.get(a.classification, 0) + 1
    else:
        avg_quality = 0
        avg_confidence = 0
        avg_spoilage = 0
        classification_counts = {}

    story.append(Paragraph("Summary Statistics", style_section))
    summary_data = [
        [Paragraph("<b>Total Analyses</b>", style_body), Paragraph(str(total), style_body_bold)],
        [Paragraph("<b>Average Quality Score</b>", style_body), Paragraph(f"{avg_quality} / 100", style_body_bold)],
        [Paragraph("<b>Average Confidence</b>", style_body), Paragraph(f"{avg_confidence}%", style_body_bold)],
        [Paragraph("<b>Average Spoilage Probability</b>", style_body), Paragraph(f"{avg_spoilage}%", style_body_bold)],
    ]
    summary_table = Table(summary_data, colWidths=[3 * inch, 2 * inch])
    summary_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), BG_LIGHT), ("BOX", (0, 0), (-1, -1), 0.5, BORDER_COLOR), ("INNERGRID", (0, 0), (-1, -1), 0.5, BORDER_COLOR), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5), ("LEFTPADDING", (0, 0), (-1, -1), 8)]))
    story.append(summary_table)
    story.append(Spacer(1, 14))

    # Classification distribution
    if classification_counts:
        story.append(Paragraph("Classification Distribution", style_section))
        dist_data = [[Paragraph("<b>Classification</b>", style_body_bold), Paragraph("<b>Count</b>", style_body_bold)]]
        for cls, count in sorted(classification_counts.items()):
            dist_data.append([Paragraph(cls, style_body), Paragraph(str(count), style_body_bold)])
        dist_table = Table(dist_data, colWidths=[3 * inch, 2 * inch])
        dist_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), BG_LIGHT), ("LINEBELOW", (0, 0), (-1, 0), 1, PRIMARY), ("BOX", (0, 0), (-1, -1), 0.5, BORDER_COLOR), ("INNERGRID", (0, 0), (-1, -1), 0.5, BORDER_COLOR), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5), ("LEFTPADDING", (0, 0), (-1, -1), 8)]))
        story.append(dist_table)
        story.append(Spacer(1, 14))

    # Analysis details table
    story.append(Paragraph("Analysis Details", style_section))
    if analyses:
        table_data = [[Paragraph("<b>Item</b>", style_body_bold), Paragraph("<b>Classification</b>", style_body_bold), Paragraph("<b>Quality</b>", style_body_bold), Paragraph("<b>Confidence</b>", style_body_bold), Paragraph("<b>Spoilage Risk</b>", style_body_bold), Paragraph("<b>Date</b>", style_body_bold)]]
        for a in analyses[:20]:
            table_data.append([
                Paragraph(str(a.item_name or "N/A"), style_body),
                Paragraph(a.classification, style_body),
                Paragraph(f"{a.quality_score}/100", style_body),
                Paragraph(f"{a.confidence * 100:.1f}%", style_body),
                Paragraph(f"{a.spoilage_probability * 100:.1f}%", style_body),
                Paragraph(a.created_at.strftime("%Y-%m-%d") if a.created_at else "N/A", style_body),
            ])
        detail_table = Table(table_data, colWidths=[1.3 * inch, 1.1 * inch, 0.8 * inch, 0.9 * inch, 1 * inch, 0.9 * inch])
        detail_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), BG_LIGHT), ("LINEBELOW", (0, 0), (-1, 0), 1, PRIMARY), ("BOX", (0, 0), (-1, -1), 0.5, BORDER_COLOR), ("INNERGRID", (0, 0), (-1, -1), 0.5, BORDER_COLOR), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
        story.append(detail_table)
    else:
        story.append(Paragraph("No freshness analysis data available.", style_body))

    story.append(Spacer(1, 14))
    story.append(HRFlowable(width="100%", thickness=0.5, color=BORDER_COLOR, spaceBefore=0, spaceAfter=8))
    story.append(Paragraph("<b>Disclaimer:</b> This report is generated from actual M2 freshness analysis data. It is intended for monitoring and quality tracking purposes.", style_disclaimer))

    doc.build(story)
    pdf_bytes = buffer.getvalue()
    buffer.close()
    return pdf_bytes


def generate_freshness_excel(analyses: List[Any]) -> bytes:
    """Generate an Excel file with freshness analysis data."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = "Freshness Analyses"

    ws.merge_cells("A1:H1")
    ws["A1"] = "FreshGuard — Freshness Summary Report"
    ws["A1"].font = Font(size=14, bold=True, color="059669")
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:H2")
    ws["A2"] = f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')}"
    ws["A2"].font = Font(size=9, color="64748b")
    ws["A2"].alignment = Alignment(horizontal="center")

    headers = ["ID", "Item Name", "Classification", "Confidence", "Quality Score", "Spoilage Probability", "Predicted Class", "Analysis Date"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    for row_idx, a in enumerate(analyses, 5):
        ws.cell(row=row_idx, column=1, value=a.id)
        ws.cell(row=row_idx, column=2, value=a.item_name or "N/A")
        ws.cell(row=row_idx, column=3, value=a.classification)
        ws.cell(row=row_idx, column=4, value=a.confidence)
        ws.cell(row=row_idx, column=5, value=a.quality_score)
        ws.cell(row=row_idx, column=6, value=a.spoilage_probability)
        ws.cell(row=row_idx, column=7, value=a.predicted_class or "N/A")
        ws.cell(row=row_idx, column=8, value=a.created_at.strftime("%Y-%m-%d %H:%M:%S") if a.created_at else "N/A")

    for col in ws.columns:
        max_length = 0
        first_cell = col[0]
        if not hasattr(first_cell, "column_letter"):
            continue
        column = first_cell.column_letter
        for cell in col:
            try:
                if cell.value and len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:
                pass
        ws.column_dimensions[column].width = min(max_length + 2, 30)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
