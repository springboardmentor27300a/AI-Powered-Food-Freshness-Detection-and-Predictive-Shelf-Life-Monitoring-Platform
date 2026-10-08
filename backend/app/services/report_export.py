"""
Reports & export service (Milestone 3 & 4).

Builds all five report types as real downloadable PDF (ReportLab) and Excel (openpyxl) files:
  1. Freshness Reports (freshness)
  2. Shelf-Life Reports (shelf_life)
  3. Inventory Quality Reports (inventory_quality)
  4. Waste Reduction Reports (waste_reduction)
  5. Storage Compliance Reports (storage_compliance)

Supports both inventory-wide datasets and per-batch drilldowns.
"""
import io
from datetime import date
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.models import FoodBatch, ImageAnalysis
from app.routers.batches import _visible_batches_query
from app.services.analytics_service import (
    expiry_analysis,
    freshness_distribution,
    inventory_insights,
    risk_analytics,
    shelf_life_analytics,
    storage_analytics,
)
from app.services.batch_insight import (
    build_batch_insight,
    prediction_for_batch,
    same_product_batches,
    score_for_batch,
    storage_for_batch,
)
from app.services.storage_service import StorageAnalyzer
from app.utils.freshness import days_to_expiry

_REPORT_TYPES = (
    "freshness",
    "shelf_life",
    "inventory_quality",
    "waste_reduction",
    "storage_compliance",
)

MONTH_NAMES = (
    "Jan", "Feb", "Mar", "Apr", "May", "Jun",
    "Jul", "Aug", "Sep", "Oct", "Nov", "Dec",
)


def is_valid_report_type(report_type: str) -> bool:
    return report_type in _REPORT_TYPES


def _fmt_date(d: Optional[date]) -> str:
    if not d:
        return "-"
    try:
        return f"{d.day:02d} {MONTH_NAMES[d.month - 1]} {d.year}"
    except Exception:
        return str(d)


def _fmt_iso(d: Optional[date]) -> str:
    if not d:
        return "-"
    try:
        if hasattr(d, "isoformat"):
            return d.isoformat()
        return str(d)
    except Exception:
        return "-"


def _status_value(status) -> str:
    if not status:
        return "-"
    value = getattr(status, "value", None)
    recommended = getattr(status, "recommended", "")
    st = getattr(status, "status", "unknown")
    if value is None:
        base = st.replace("_", " ").capitalize()
    else:
        base = str(value)
    if recommended:
        return f"{base} (rec: {recommended}) [{st}]"
    return f"{base} [{st}]"


# ---------------------------------------------------------------------------
# Per-batch report data
# ---------------------------------------------------------------------------
def batch_report_data(db, batch: FoodBatch) -> dict:
    insight = build_batch_insight(db, batch, same_product_batches(db, batch))
    pred = insight["prediction"]
    storage = insight["storage"]
    score = insight["score"]

    details = [
        ("Batch ID", batch.batch_id),
        ("Food Name", batch.food_name),
        ("Category", batch.category or "Unknown"),
        ("Packaging Type", batch.packaging_type),
        ("Storage Location", batch.storage_location),
        ("Received Date", _fmt_date(batch.received_date)),
        ("Registered Expiry", _fmt_date(batch.expiry_date)),
        ("Days to Expiry", str(days_to_expiry(batch.expiry_date))),
    ]
    pred_rows = [
        ("Estimated Remaining Shelf Life", f"{pred.estimated_remaining_days} day(s)"),
        ("Expected Expiry Date", _fmt_date(pred.expected_expiry_date)),
        ("Calendar Remaining Days", str(pred.calendar_remaining_days)),
        ("Freshness Score (visual proxy)", f"{pred.freshness_score}"),
        ("Shelf-Life Score", f"{pred.shelf_life_score}"),
        ("Storage Condition Score", f"{pred.storage_condition_score}"),
        ("Spoilage Risk", pred.spoilage_risk),
        ("Risk Score", str(pred.risk_score)),
        ("Storage Delta Days", f"{pred.storage_delta_days:+d}"),
    ]
    storage_rows = [
        (name, _status_value(getattr(storage, field, None)), getattr(getattr(storage, field, None), "status", "-"))
        for name, field in (
            ("Temperature", "temperature"),
            ("Humidity", "humidity"),
            ("Air Circulation", "air_circulation"),
            ("Light Exposure", "light_exposure"),
            ("Storage Duration", "duration"),
        )
    ]
    score_rows = [
        ("Visual Condition (40%)", str(score.visual_condition_score)),
        ("Storage Conditions (25%)", str(score.storage_condition_score)),
        ("Shelf-Life Prediction (20%)", str(score.shelf_life_score)),
        ("Product Age (15%)", str(score.product_age_score)),
        ("Overall Freshness Score", str(score.overall_score)),
        ("Freshness Status", score.freshness_status),
    ]
    recs = [
        {"category": r.category, "message": r.message, "priority": r.priority}
        for r in insight.get("recommendations", [])
    ]

    return {
        "title": f"Batch Freshness Report - {batch.batch_id}",
        "report_type": "batch",
        "batch_id": batch.batch_id,
        "generated_on": _fmt_date(date.today()),
        "details": details,
        "prediction": pred_rows,
        "storage": storage_rows,
        "score": score_rows,
        "factors": pred.factors,
        "storage_recommendations": storage.optimization_recommendations,
        "recommendations": recs,
    }


# ---------------------------------------------------------------------------
# Inventory report data (Aggregated)
# ---------------------------------------------------------------------------
def inventory_report_data(db, user) -> dict:
    today = date.today()
    fmt = _fmt_date(today)

    insights = inventory_insights(db, user)
    expiry = expiry_analysis(db, user)
    shelf = shelf_life_analytics(db, user)
    storage = storage_analytics(db, user)
    risk = risk_analytics(db, user)
    dist = freshness_distribution(db, user)

    summary = [
        ("Total Products", str(insights["total_products"])),
        ("Fresh", str(insights["fresh_count"])),
        ("Acceptable", str(insights["acceptable_count"])),
        ("Needs Attention", str(insights["needs_attention_count"])),
        ("Spoiled", str(insights["spoiled_count"])),
        ("At Risk", str(insights["at_risk_count"])),
        ("Expired", str(insights["expired_count"])),
        ("Expiring within 7 days", str(insights["expiring_within_7_count"])),
        ("Average Freshness Score", str(insights["average_freshness_score"])),
        ("Average Remaining Shelf Life (days)", str(insights["average_shelf_life_days"])),
        ("Average Storage Compliance", str(insights["average_storage_compliance"])),
    ]
    expiry_rows = [
        ("Expired", str(expiry["expired"])),
        ("Expiring today", str(expiry["today"])),
        ("Expiring within 1 day", str(expiry["within_1_day"])),
        ("Expiring within 3 days", str(expiry["within_3_days"])),
        ("Expiring within 7 days", str(expiry["within_7_days"])),
    ]
    storage_rows = [
        ("Overall Storage Compliance", f"{storage['overall_compliance']}%"),
        ("Monitored Batches", str(storage["monitored_batches"])),
        (
            "Temperature - Good / Warning / Critical",
            f"{storage['by_parameter']['temperature']['good']} / "
            f"{storage['by_parameter']['temperature']['warning']} / "
            f"{storage['by_parameter']['temperature']['critical']}",
        ),
        (
            "Humidity - Good / Warning / Critical",
            f"{storage['by_parameter']['humidity']['good']} / "
            f"{storage['by_parameter']['humidity']['warning']} / "
            f"{storage['by_parameter']['humidity']['critical']}",
        ),
    ]
    risk_rows = [
        ("Low", str(risk["low"])),
        ("Moderate", str(risk["moderate"])),
        ("High", str(risk["high"])),
        ("Critical", str(risk["critical"])),
    ]

    return {
        "title": "Inventory Quality Report",
        "report_type": "inventory_quality",
        "generated_on": fmt,
        "batch_id": "All Inventory",
        "summary": summary,
        "expiry": expiry_rows,
        "storage": storage_rows,
        "risk": risk_rows,
        "distribution": [
            ("Fresh", str(dist.get("Fresh", 0))),
            ("Acceptable", str(dist.get("Acceptable", 0))),
            ("Needs Attention", str(dist.get("Needs Attention", 0))),
            ("Spoiled", str(dist.get("Spoiled", 0))),
        ],
        "shelf_life": [
            ("Average remaining days", str(shelf["average_remaining_days"])),
            ("Expiring in 1 day", str(shelf["expiring_in_1_day"])),
            ("Expiring in 3 days", str(shelf["expiring_in_3_days"])),
            ("Expiring in 7 days", str(shelf["expiring_in_7_days"])),
        ],
        "waste_risk": insights.get("waste_risk_items", []),
        "details": summary,
    }


# ---------------------------------------------------------------------------
# Full table reports for each of the 5 Milestone 4 report types
# ---------------------------------------------------------------------------
def table_report_data(db, user, report_type: str) -> dict:
    today = date.today()
    fmt = _fmt_date(today)

    batches = db.scalars(_visible_batches_query(db, user).order_by(FoodBatch.expiry_date.asc())).all()

    # 1. Freshness Report
    if report_type == "freshness":
        headers = [
            "Food Item", "Category", "Batch ID", "Freshness Score",
            "Freshness Category", "Spoilage Probability", "Analysis Date", "Confidence", "Status"
        ]
        table_rows = []
        items = []

        fresh_count = 0
        good_count = 0
        acceptable_count = 0
        near_spoil_count = 0
        spoiled_count = 0
        total_score = 0.0

        for b in batches:
            # Check if there is an image analysis linked to this batch
            latest_img = (
                db.query(ImageAnalysis)
                .filter(ImageAnalysis.batch_id_ref == b.batch_id)
                .order_by(ImageAnalysis.created_at.desc())
                .first()
            )
            pred = prediction_for_batch(db, b)
            sc = score_for_batch(b, pred)

            if latest_img:
                score_val = round(latest_img.confidence_score * 100, 1)
                category_val = latest_img.classification
                prob_val = f"{latest_img.spoilage_probability * 100:.1f}%"
                conf_val = f"{latest_img.confidence_score * 100:.0f}%"
                analysis_date = _fmt_date(latest_img.created_at.date() if latest_img.created_at else today)
            else:
                score_val = round(sc.overall_score, 1)
                category_val = sc.freshness_status
                prob_val = f"{pred.risk_score:.1f}%" if pred.risk_score else "10.0%"
                conf_val = "85%"
                analysis_date = _fmt_date(b.received_date)

            status_val = b.freshness_status
            total_score += score_val

            if category_val == "Fresh":
                fresh_count += 1
            elif category_val == "Good":
                good_count += 1
            elif category_val == "Acceptable":
                acceptable_count += 1
            elif category_val == "Near Spoilage":
                near_spoil_count += 1
            else:
                spoiled_count += 1

            row = [
                b.food_name,
                b.category or "-",
                b.batch_id,
                f"{score_val}/100",
                category_val,
                prob_val,
                analysis_date,
                conf_val,
                status_val,
            ]
            table_rows.append(row)
            items.append({
                "food_name": b.food_name,
                "category": b.category or "-",
                "batch_id": b.batch_id,
                "freshness_score": score_val,
                "freshness_category": category_val,
                "spoilage_probability": prob_val,
                "analysis_date": analysis_date,
                "confidence": conf_val,
                "status": status_val,
            })

        avg_score = round(total_score / max(len(batches), 1), 1)
        summary = [
            ("Total Batches Analyzed", str(len(batches))),
            ("Average Freshness Score", f"{avg_score}/100"),
            ("Fresh", str(fresh_count)),
            ("Good", str(good_count)),
            ("Acceptable", str(acceptable_count)),
            ("Near Spoilage", str(near_spoil_count)),
            ("Spoiled", str(spoiled_count)),
        ]

        return {
            "title": "Freshness Assessment Report",
            "report_type": "freshness",
            "generated_on": fmt,
            "batch_id": "All Inventory",
            "summary": summary,
            "table_header": headers,
            "table_rows": table_rows,
            "details": table_rows,
            "items": items,
        }

    # 2. Shelf-Life Report
    if report_type == "shelf_life":
        headers = [
            "Product", "Batch ID", "Storage Duration", "Estimated Remaining",
            "Expiry Date", "Risk Level", "Confidence"
        ]
        table_rows = []
        items = []

        expiring_soon = 0
        expired = 0
        total_remaining = 0

        for b in batches:
            pred = prediction_for_batch(db, b)
            duration_days = max(0, (today - b.received_date).days)
            remaining_days = pred.estimated_remaining_days
            total_remaining += remaining_days

            if remaining_days <= 0:
                expired += 1
            elif remaining_days <= 3:
                expiring_soon += 1

            conf_str = "90%" if pred.shelf_life_score >= 80 else ("80%" if pred.shelf_life_score >= 50 else "70%")

            row = [
                b.food_name,
                b.batch_id,
                f"{duration_days} days",
                f"{remaining_days} days",
                _fmt_iso(pred.expected_expiry_date),
                pred.spoilage_risk,
                conf_str,
            ]
            table_rows.append(row)
            items.append({
                "food_name": b.food_name,
                "batch_id": b.batch_id,
                "storage_duration": f"{duration_days} days",
                "remaining_days": remaining_days,
                "expected_expiry_date": _fmt_iso(pred.expected_expiry_date),
                "spoilage_risk": pred.spoilage_risk,
                "confidence": conf_str,
            })

        avg_remaining = round(total_remaining / max(len(batches), 1), 1)
        summary = [
            ("Total Monitored Products", str(len(batches))),
            ("Average Remaining Shelf Life", f"{avg_remaining} days"),
            ("Expiring Soon (<= 3 days)", str(expiring_soon)),
            ("Expired Products", str(expired)),
        ]

        return {
            "title": "Shelf-Life Prediction Report",
            "report_type": "shelf_life",
            "generated_on": fmt,
            "batch_id": "All Inventory",
            "summary": summary,
            "table_header": headers,
            "table_rows": table_rows,
            "details": table_rows,
            "items": items,
        }

    # 3. Inventory Quality Report
    if report_type == "inventory_quality":
        headers = [
            "Batch ID", "Food Name", "Category", "Available Qty",
            "Unit", "Expiry Date", "Freshness Status", "Quality Score", "Risk Level"
        ]
        table_rows = []
        items = []

        iv = inventory_insights(db, user)
        for b in batches:
            pred = prediction_for_batch(db, b)
            sc = score_for_batch(b, pred)
            row = [
                b.batch_id,
                b.food_name,
                b.category or "-",
                str(b.available_quantity),
                b.unit,
                _fmt_iso(b.expiry_date),
                sc.freshness_status,
                f"{sc.overall_score}/100",
                pred.spoilage_risk,
            ]
            table_rows.append(row)
            items.append({
                "batch_id": b.batch_id,
                "food_name": b.food_name,
                "category": b.category or "-",
                "available_quantity": b.available_quantity,
                "unit": b.unit,
                "expiry_date": _fmt_iso(b.expiry_date),
                "freshness_status": sc.freshness_status,
                "overall_score": sc.overall_score,
                "spoilage_risk": pred.spoilage_risk,
            })

        summary = [
            ("Total Inventory Items", str(iv["total_products"])),
            ("Fresh", str(iv["fresh_count"])),
            ("Acceptable", str(iv["acceptable_count"])),
            ("Needs Attention", str(iv["needs_attention_count"])),
            ("Spoiled", str(iv["spoiled_count"])),
            ("At Risk Items", str(iv["at_risk_count"])),
            ("Average Freshness Score", str(iv["average_freshness_score"])),
        ]

        return {
            "title": "Inventory Quality Report",
            "report_type": "inventory_quality",
            "generated_on": fmt,
            "batch_id": "All Inventory",
            "summary": summary,
            "table_header": headers,
            "table_rows": table_rows,
            "details": table_rows,
            "items": items,
        }

    # 4. Waste Reduction Report
    if report_type == "waste_reduction":
        headers = [
            "Batch ID", "Product", "Category", "Available Qty",
            "Remaining Days", "Freshness Status", "Spoilage Risk", "Recommended Action", "Rotation Suggestion"
        ]
        table_rows = []
        items = []

        iv = inventory_insights(db, user)
        waste_items = iv.get("waste_risk_items", [])
        waste_batch_ids = {w["batch_id"] for w in waste_items}

        for b in batches:
            pred = prediction_for_batch(db, b)
            sc = score_for_batch(b, pred)

            # Include at-risk, near-spoilage, or expiring batches
            is_waste_risk = (
                b.batch_id in waste_batch_ids
                or pred.estimated_remaining_days <= 3
                or pred.spoilage_risk in ("High", "Critical")
                or sc.freshness_status in ("Needs Attention", "Spoiled")
            )
            if not is_waste_risk:
                continue

            if pred.estimated_remaining_days <= 0:
                rec_action = "Discard or compost safely. Do not consume."
                rotation = "Remove from active inventory"
            elif pred.estimated_remaining_days <= 1:
                rec_action = "Priority consumption or immediate markdown sale."
                rotation = "Immediate FEFO release"
            elif pred.estimated_remaining_days <= 3:
                rec_action = "Front-display promotion and expedited sale."
                rotation = "First-Expiring First-Out"
            else:
                rec_action = "Inspect storage conditions; consume promptly."
                rotation = "Standard FIFO rotation"

            row = [
                b.batch_id,
                b.food_name,
                b.category or "-",
                f"{b.available_quantity} {b.unit}",
                f"{pred.estimated_remaining_days} days",
                sc.freshness_status,
                pred.spoilage_risk,
                rec_action,
                rotation,
            ]
            table_rows.append(row)
            items.append({
                "batch_id": b.batch_id,
                "food_name": b.food_name,
                "category": b.category or "-",
                "available_quantity": f"{b.available_quantity} {b.unit}",
                "remaining_days": pred.estimated_remaining_days,
                "freshness_status": sc.freshness_status,
                "spoilage_risk": pred.spoilage_risk,
                "recommended_action": rec_action,
                "rotation": rotation,
            })

        summary = [
            ("Total Waste-Risk Batches", str(len(table_rows))),
            ("Critical / High Risk", str(iv["at_risk_count"])),
            ("Expiring <= 7 Days", str(iv["expiring_within_7_count"])),
            ("Expired Batches", str(iv["expired_count"])),
        ]

        return {
            "title": "Waste Reduction & Risk Report",
            "report_type": "waste_reduction",
            "generated_on": fmt,
            "batch_id": "All Inventory",
            "summary": summary,
            "table_header": headers,
            "table_rows": table_rows,
            "details": table_rows,
            "items": items,
        }

    # 5. Storage Compliance Report
    if report_type == "storage_compliance":
        headers = [
            "Product", "Batch ID", "Current Temp", "Recommended Temp",
            "Current Humidity", "Recommended Humidity", "Air Circulation", "Light Exposure",
            "Compliance", "Risk Level", "Recommendation"
        ]
        table_rows = []
        items = []

        sa = StorageAnalyzer()
        good_count = 0
        warn_count = 0
        crit_count = 0

        for b in batches:
            a = sa.analyze(
                batch_id=b.batch_id,
                food_name=b.food_name,
                category=b.category,
                received_date=b.received_date,
                temperature_c=b.temperature_c,
                humidity_pct=b.humidity_pct,
                air_circulation=b.air_circulation,
                light_exposure=b.light_exposure,
                packaging_type=b.packaging_type,
                storage_location=b.storage_location,
            )
            pred = prediction_for_batch(db, b)

            if a.compliance_score >= 80:
                good_count += 1
            elif a.compliance_score >= 50:
                warn_count += 1
            else:
                crit_count += 1

            first_rec = a.optimization_recommendations[0] if a.optimization_recommendations else "Maintain current storage conditions."

            row = [
                b.food_name,
                b.batch_id,
                f"{b.temperature_c}°C" if b.temperature_c is not None else "Not set",
                a.temperature.recommended,
                f"{b.humidity_pct}%" if b.humidity_pct is not None else "Not set",
                a.humidity.recommended,
                a.air_circulation.status.capitalize(),
                a.light_exposure.status.capitalize(),
                f"{a.compliance_score}% ({a.condition_status})",
                pred.spoilage_risk,
                first_rec,
            ]
            table_rows.append(row)
            items.append({
                "food_name": b.food_name,
                "batch_id": b.batch_id,
                "current_temp": f"{b.temperature_c}°C" if b.temperature_c is not None else "Not set",
                "recommended_temp": a.temperature.recommended,
                "current_humidity": f"{b.humidity_pct}%" if b.humidity_pct is not None else "Not set",
                "recommended_humidity": a.humidity.recommended,
                "air_circulation": a.air_circulation.status.capitalize(),
                "light_exposure": a.light_exposure.status.capitalize(),
                "compliance": f"{a.compliance_score}%",
                "condition_status": a.condition_status,
                "risk_level": pred.spoilage_risk,
                "recommendation": first_rec,
            })

        summary = [
            ("Total Batches Monitored", str(len(batches))),
            ("Fully Compliant (>= 80%)", str(good_count)),
            ("Needs Attention (50-79%)", str(warn_count)),
            ("Non-Compliant (< 50%)", str(crit_count)),
        ]

        return {
            "title": "Storage Condition & Compliance Report",
            "report_type": "storage_compliance",
            "generated_on": fmt,
            "batch_id": "All Inventory",
            "summary": summary,
            "table_header": headers,
            "table_rows": table_rows,
            "details": table_rows,
            "items": items,
        }

    # Fallback to inventory report data
    return inventory_report_data(db, user)


# ---------------------------------------------------------------------------
# PDF Generation (Landscape for tables with > 5 columns)
# ---------------------------------------------------------------------------
def build_pdf(data: dict) -> bytes:
    buf = io.BytesIO()

    has_wide_table = bool(data.get("table_header") and len(data.get("table_header", [])) > 5)
    page_size = landscape(A4) if has_wide_table else A4

    doc = SimpleDocTemplate(
        buf,
        pagesize=page_size,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        title=f"{data.get('title', 'Food Freshness Report')}",
    )

    style_h = ParagraphStyle(
        "h",
        fontName="Helvetica-Bold",
        fontSize=15,
        leading=18,
        textColor=colors.HexColor("#14532d"),
    )
    style_sub = ParagraphStyle(
        "sub",
        fontName="Helvetica",
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#64748b"),
    )
    style_s = ParagraphStyle(
        "sec",
        fontName="Helvetica-Bold",
        fontSize=10.5,
        leading=14,
        spaceBefore=7,
        spaceAfter=3,
        textColor=colors.HexColor("#166534"),
    )
    style_cell = ParagraphStyle(
        "cell",
        fontName="Helvetica",
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor("#1e293b"),
    )
    style_head_cell = ParagraphStyle(
        "head_cell",
        fontName="Helvetica-Bold",
        fontSize=8,
        leading=10,
        textColor=colors.white,
    )
    style_disclaimer = ParagraphStyle(
        "disc",
        fontName="Helvetica",
        fontSize=7,
        leading=9,
        textColor=colors.HexColor("#64748b"),
    )

    story = [
        Paragraph(str(data.get("title", "Food Freshness Report")), style_h),
        Paragraph(
            f"Scope: {data.get('batch_id', 'All Inventory')} &nbsp;|&nbsp; "
            f"Generated On: {data.get('generated_on', '-')} &nbsp;|&nbsp; "
            f"Infosys Springboard - Food Freshness Monitoring Platform",
            style_sub,
        ),
        Spacer(1, 5),
    ]

    # Render summary cards/table if present
    summary_rows = data.get("summary")
    if summary_rows:
        story.append(Paragraph("Executive Summary", style_s))
        s_data = []
        for i in range(0, len(summary_rows), 2):
            row = []
            row.append(summary_rows[i][0])
            row.append(summary_rows[i][1])
            if i + 1 < len(summary_rows):
                row.append(summary_rows[i + 1][0])
                row.append(summary_rows[i + 1][1])
            else:
                row.extend(["", ""])
            s_data.append(row)

        s_table = Table(s_data)
        s_table.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f1f5f9")),
            ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#f1f5f9")),
            ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
            ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
            ("TOPPADDING", (0, 0), (-1, -1), 2.5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
        ]))
        story.append(s_table)
        story.append(Spacer(1, 6))

    # Render tabular report if table_header exists
    table_header = data.get("table_header")
    table_rows = data.get("table_rows") or []

    if table_header and table_rows:
        story.append(Paragraph("Detailed Records", style_s))

        # Format header with paragraphs
        fmt_header = [Paragraph(str(h), style_head_cell) for h in table_header]
        fmt_rows = [fmt_header]

        for r in table_rows:
            fmt_row = [Paragraph(str(c), style_cell) for c in r]
            fmt_rows.append(fmt_row)

        t = Table(fmt_rows, repeatRows=1)
        t_style = [
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#cbd5e1")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#166534")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 2.5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ]
        # Alternate row backgrounds
        for row_idx in range(1, len(fmt_rows)):
            bg = colors.HexColor("#ffffff") if row_idx % 2 == 1 else colors.HexColor("#f8fafc")
            t_style.append(("BACKGROUND", (0, row_idx), (-1, row_idx), bg))

        t.setStyle(TableStyle(t_style))
        story.append(t)
        story.append(Spacer(1, 8))

    elif not table_header:
        # Standard per-batch structured sections
        def add_2col_section(title, rows):
            if not rows:
                return
            story.append(Paragraph(title, style_s))
            formatted = [[Paragraph(str(k), style_cell), Paragraph(str(v), style_cell)] for k, v in rows]
            t = Table(formatted, colWidths=[60 * mm, None])
            t.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#e2e8f0")),
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f0fdf4")),
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("TOPPADDING", (0, 0), (-1, -1), 2.5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
            ]))
            story.append(t)
            story.append(Spacer(1, 5))

        title_map = {
            "details": "Food & Batch Information",
            "prediction": "Shelf-Life Prediction",
            "storage": "Storage Conditions",
            "score": "Overall Freshness Score Breakdown",
            "distribution": "Freshness Distribution",
            "shelf_life": "Shelf-Life Analytics",
        }
        for sec, title in title_map.items():
            add_2col_section(title, data.get(sec) or [])

        recs = data.get("recommendations") or []
        if recs:
            story.append(Paragraph("Recommendations", style_s))
            for r in recs[:25]:
                story.append(Paragraph(f"• [{r.get('category', '').upper()}] {r.get('message', '')}", style_cell))
            story.append(Spacer(1, 5))

    # Disclaimer
    story.append(Spacer(1, 6))
    story.append(Paragraph(
        "Disclaimer: This report is generated by the AI-powered Food Freshness Monitoring Platform. "
        "Freshness scores and shelf-life predictions are computer-vision and rule-assisted estimates "
        "and should be supplemented with standard food safety guidelines.",
        style_disclaimer,
    ))

    doc.build(story)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Excel Generation
# ---------------------------------------------------------------------------
def build_excel(data: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = (data.get("report_type") or "Report")[:30]

    # Styles
    title_font = Font(name="Calibri", size=14, bold=True, color="14532D")
    meta_font = Font(name="Calibri", size=9, italic=True, color="475569")
    sec_font = Font(name="Calibri", size=11, bold=True, color="166534")
    head_font = Font(name="Calibri", size=9.5, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="166534")
    zebra_fill = PatternFill("solid", fgColor="F8FAFC")
    border_side = Side(style="thin", color="CBD5E1")
    cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    cell_font = Font(name="Calibri", size=9)

    # Title & Metadata
    ws.append([data.get("title", "Food Freshness Report")])
    ws.cell(ws.max_row, 1).font = title_font

    ws.append([
        f"Scope: {data.get('batch_id', 'All Inventory')}",
        f"Generated: {data.get('generated_on', '-')}",
        "Infosys Springboard - Food Freshness Platform"
    ])
    for col_idx in range(1, 4):
        ws.cell(ws.max_row, col_idx).font = meta_font

    ws.append([])

    # Executive Summary if present
    summary_rows = data.get("summary")
    if summary_rows:
        ws.append(["Executive Summary"])
        ws.cell(ws.max_row, 1).font = sec_font

        for label, val in summary_rows:
            ws.append([label, val])
            row_idx = ws.max_row
            ws.cell(row_idx, 1).font = Font(bold=True, size=9)
            ws.cell(row_idx, 1).border = cell_border
            ws.cell(row_idx, 2).font = cell_font
            ws.cell(row_idx, 2).border = cell_border

        ws.append([])

    # Detailed Table
    table_header = data.get("table_header")
    table_rows = data.get("table_rows") or []

    if table_header and table_rows:
        ws.append(["Detailed Records"])
        ws.cell(ws.max_row, 1).font = sec_font

        # Header row
        ws.append(list(table_header))
        h_row_idx = ws.max_row
        for col_idx in range(1, len(table_header) + 1):
            cell = ws.cell(h_row_idx, col_idx)
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = cell_border

        # Data rows
        for r_idx, r in enumerate(table_rows):
            ws.append(list(r))
            curr_row = ws.max_row
            for col_idx in range(1, len(r) + 1):
                cell = ws.cell(curr_row, col_idx)
                cell.font = cell_font
                cell.border = cell_border
                if r_idx % 2 == 1:
                    cell.fill = zebra_fill
    elif not table_header:
        # Key-value fallback
        details = data.get("details") or []
        if details:
            ws.append(["Batch Details"])
            ws.cell(ws.max_row, 1).font = sec_font
            for row in details:
                if isinstance(row, (list, tuple)) and len(row) >= 2:
                    ws.append([row[0], row[1]])
                    ws.cell(ws.max_row, 1).font = Font(bold=True, size=9)
                    ws.cell(ws.max_row, 1).border = cell_border
                    ws.cell(ws.max_row, 2).font = cell_font
                    ws.cell(ws.max_row, 2).border = cell_border

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(min(max_len + 3, 40), 12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
